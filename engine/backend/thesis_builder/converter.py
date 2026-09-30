"""Markdown→DOCX 转换引擎。

流程：
1. 解析 Markdown（parser.parse_markdown），python-docx 按格式规范逐段构建
2. 合并模板封面/封底（template_merge）
3. 添加页眉页脚、分节符、脚注、目录域（postprocess）
4. Word 分页后处理：更新目录域 + 补奇偶页空白填充（word_pass.ps1）
"""

from __future__ import annotations

import re
import subprocess
from pathlib import Path
from typing import Dict, List, Optional

from docx import Document
from docx.shared import Pt, Cm, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from lxml import etree

from ..process_utils import hidden_process_kwargs
from .parser import ParsedDocument, ParsedSection, TableData, parse_markdown
from .styles import (
    ThesisConfig, PageConfig, FontConfig, SpacingConfig,
    FONT_SIMSUN, FONT_SIMHEI, FONT_FANGSONG, FONT_TIMES,
)
from .template_merge import load_template_profile, merge_template
from .postprocess import (
    setup_headers_footers, add_section_break, add_footnotes,
    add_toc_field,
)


# ═══════════════════════════════════════════
# 段落构建辅助函数
# ═══════════════════════════════════════════

def _set_run_font(run, east_asian: str, latin: str, size_pt: float,
                   bold: bool = False, italic: bool = False) -> None:
    """设置 run 的字体属性"""
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    # 设置西文字体
    run.font.name = latin
    # 设置东亚字体
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rpr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), east_asian)
    rFonts.set(qn('w:ascii'), latin)
    rFonts.set(qn('w:hAnsi'), latin)


def _set_paragraph_spacing(paragraph, before_pt: float = 0, after_pt: float = 0,
                            line_spacing: float = 1.25, first_line_chars: int = 0,
                            hanging_chars: int = 0) -> None:
    """设置段落间距和缩进"""
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before_pt)
    pf.space_after = Pt(after_pt)
    pf.line_spacing = line_spacing

    if first_line_chars > 0:
        # 首行缩进：按字符数 × 字号
        font_size = paragraph.runs[0].font.size if paragraph.runs else Pt(12)
        pf.first_line_indent = int(font_size * first_line_chars)
    elif hanging_chars > 0:
        font_size = paragraph.runs[0].font.size if paragraph.runs else Pt(12)
        pf.first_line_indent = -int(font_size * hanging_chars)


def _add_paragraph(doc: Document, text: str, font_cfg: FontConfig,
                    spacing_cfg: SpacingConfig, alignment=None,
                    east_asian: str = None, latin: str = None,
                    size_pt: float = None, bold: bool = False,
                    italic: bool = False, first_line_chars: int = 0,
                    hanging_chars: int = 0,
                    before_pt: float = 0, after_pt: float = 0) -> 'Paragraph':
    """添加一个格式化段落"""
    para = doc.add_paragraph()
    if alignment is not None:
        para.alignment = alignment
    run = para.add_run(text)
    _set_run_font(
        run,
        east_asian or font_cfg.body_east_asian,
        latin or font_cfg.body_latin,
        size_pt or font_cfg.body_size_pt,
        bold=bold, italic=italic,
    )
    _set_paragraph_spacing(
        para,
        before_pt=before_pt, after_pt=after_pt,
        line_spacing=spacing_cfg.line_spacing,
        first_line_chars=first_line_chars,
        hanging_chars=hanging_chars,
    )
    return para


def _add_body_paragraph(doc: Document, text: str, cfg: ThesisConfig) -> 'Paragraph':
    """添加正文段落（宋体小四，首行缩进2字符）"""
    return _add_paragraph(
        doc, text, cfg.fonts, cfg.spacing,
        first_line_chars=cfg.spacing.first_line_indent_chars,
    )


def _add_styled_heading(doc: Document, title: str, style_id_name: str,
                        line_spacing: float) -> 'Paragraph':
    """用段落样式添加标题（不带 run 直接格式，避免污染目录条目）"""
    para = doc.add_paragraph()
    para.style = doc.styles[style_id_name]
    para.paragraph_format.line_spacing = line_spacing
    para.add_run(title)
    return para


def _add_heading1(doc: Document, title: str, cfg: ThesisConfig) -> 'Paragraph':
    """添加一级标题（黑体三号，居中，段前段后各1行）"""
    return _add_styled_heading(doc, title, 'Heading 1', cfg.spacing.line_spacing)


def _add_heading2(doc: Document, title: str, cfg: ThesisConfig) -> 'Paragraph':
    """添加二级标题（宋体四号加粗，左顶格）"""
    return _add_styled_heading(doc, title, 'Heading 2', cfg.spacing.line_spacing)


def _add_heading3(doc: Document, title: str, cfg: ThesisConfig) -> 'Paragraph':
    """添加三级标题（宋体小四号加粗，左顶格）"""
    return _add_styled_heading(doc, title, 'Heading 3', cfg.spacing.line_spacing)


# ═══════════════════════════════════════════
# 特殊章节构建
# ═══════════════════════════════════════════

def _build_abstract_cn(doc: Document, parsed: ParsedDocument, cfg: ThesisConfig) -> None:
    """构建中文摘要"""
    # 标题：摘  要（中间空2格）
    para = _add_paragraph(
        doc, "摘  要", cfg.fonts, cfg.spacing,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        east_asian=cfg.fonts.abstract_title_font,
        size_pt=cfg.fonts.abstract_title_size_pt,
        bold=True,
        before_pt=cfg.spacing.abstract_before_pt,
        after_pt=cfg.spacing.abstract_after_pt,
    )
    # 正文（摘要不放脚注，剔除脚注标记）
    content = parsed.get_abstract_content()
    if content:
        _add_body_paragraph(doc, re.sub(r'\[\^\d+\]', '', content), cfg)
    # 关键词
    keywords = parsed.get_abstract_keywords()
    if keywords:
        kw_para = doc.add_paragraph()
        kw_para.paragraph_format.space_before = Pt(cfg.spacing.keywords_before_pt)
        kw_para.paragraph_format.line_spacing = cfg.spacing.line_spacing
        # 规范：「关键词」前空2格（按字符缩进）
        pPr = kw_para._element.get_or_add_pPr()
        pPr.append(parse_xml(
            f'<w:ind {nsdecls("w")} w:firstLineChars="200"'
            f' w:firstLine="{int(cfg.fonts.keywords_label_size_pt * 2 * 20)}"/>'
        ))
        # "关键词"标签（黑体四号）
        label_run = kw_para.add_run("关键词：")
        _set_run_font(label_run, cfg.fonts.keywords_label_font,
                      FONT_TIMES, cfg.fonts.keywords_label_size_pt, bold=True)
        # 关键词内容（宋体小四号）
        kw_run = kw_para.add_run(keywords)
        _set_run_font(kw_run, cfg.fonts.body_east_asian,
                      cfg.fonts.body_latin, cfg.fonts.body_size_pt)


def _build_abstract_en(doc: Document, parsed: ParsedDocument, cfg: ThesisConfig) -> None:
    """构建英文摘要"""
    # 标题：Abstract（Times New Roman 三号加粗）
    para = _add_paragraph(
        doc, "Abstract", cfg.fonts, cfg.spacing,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        east_asian=cfg.fonts.eng_abstract_title_font,
        latin=cfg.fonts.eng_abstract_title_font,
        size_pt=cfg.fonts.eng_abstract_title_size_pt,
        bold=True,
        before_pt=cfg.spacing.abstract_before_pt,
        after_pt=cfg.spacing.abstract_after_pt,
    )
    # 正文
    content = parsed.get_eng_abstract_content()
    if content:
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Pt(cfg.fonts.body_size_pt * 2)
        para.paragraph_format.line_spacing = cfg.spacing.line_spacing
        run = para.add_run(content)
        _set_run_font(run, FONT_SIMSUN, FONT_TIMES, cfg.fonts.body_size_pt)
    # 英文关键词
    keywords = parsed.get_eng_keywords()
    if keywords:
        kw_para = doc.add_paragraph()
        kw_para.paragraph_format.space_before = Pt(cfg.spacing.keywords_before_pt)
        kw_para.paragraph_format.line_spacing = cfg.spacing.line_spacing
        # "Key words:" 标签（Times New Roman 四号加粗，与标准文档一致）
        label_run = kw_para.add_run("Key words: ")
        _set_run_font(label_run, FONT_TIMES, FONT_TIMES,
                      cfg.fonts.keywords_label_size_pt, bold=True)
        # 关键词内容（小四号）
        kw_run = kw_para.add_run(keywords)
        _set_run_font(kw_run, FONT_SIMSUN, FONT_TIMES, cfg.fonts.body_size_pt)


def _build_toc(doc: Document, cfg: ThesisConfig) -> None:
    """构建目录页"""
    # 标题：目  录
    _add_paragraph(
        doc, "目  录", cfg.fonts, cfg.spacing,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        east_asian=FONT_SIMSUN,
        size_pt=cfg.fonts.abstract_title_size_pt,
        bold=True,
        before_pt=cfg.spacing.abstract_before_pt,
        after_pt=cfg.spacing.abstract_after_pt,
    )
    # 添加 TOC 域代码
    add_toc_field(doc)


def _build_references(doc: Document, refs: List[str], cfg: ThesisConfig) -> None:
    """构建参考文献"""
    if not refs:
        return
    for ref_no, ref_text in enumerate(refs, 1):
        # 重排编号：去除原始编号后按顺序补 [n]（与标准文档一致）
        content = re.sub(r'^\[\d+\]\s*', '', ref_text)
        para = doc.add_paragraph()
        para.paragraph_format.line_spacing = cfg.spacing.line_spacing
        # 悬挂缩进
        para.paragraph_format.first_line_indent = Pt(
            -(cfg.fonts.ref_size_pt * cfg.spacing.ref_hanging_indent_chars)
        )
        para.paragraph_format.left_indent = Pt(
            cfg.fonts.ref_size_pt * cfg.spacing.ref_hanging_indent_chars
        )

        # 编号 [n]
        num_run = para.add_run(f'[{ref_no}] ')
        _set_run_font(num_run, FONT_SIMSUN, FONT_TIMES, cfg.fonts.ref_size_pt)

        # 判断是否为外文文献
        if _is_foreign_ref(ref_text):
            m = re.match(r'^([A-Za-z][^.]*\.\s*)(.*?)(\[)', content)
            if m:
                run1 = para.add_run(m.group(1))
                _set_run_font(run1, FONT_SIMSUN, FONT_TIMES, cfg.fonts.ref_size_pt)
                run2 = para.add_run(m.group(2))
                _set_run_font(run2, FONT_TIMES, FONT_TIMES, cfg.fonts.ref_size_pt, italic=True)
                run3 = para.add_run(content[m.start(3):])
                _set_run_font(run3, FONT_SIMSUN, FONT_TIMES, cfg.fonts.ref_size_pt)
            else:
                run = para.add_run(content)
                _set_run_font(run, FONT_SIMSUN, FONT_TIMES, cfg.fonts.ref_size_pt)
        else:
            run = para.add_run(content)
            _set_run_font(run, FONT_SIMSUN, FONT_TIMES, cfg.fonts.ref_size_pt)


def _is_foreign_ref(text: str) -> bool:
    """判断是否为外文参考文献"""
    content = re.sub(r'^\[\d+\]\s*', '', text)
    m = re.match(r'^([A-Za-z][^.]*\.\s*)', content)
    if not m:
        return False
    author = m.group(1)
    latin_count = sum(1 for c in author if c.isalpha())
    return latin_count > len(author) * 0.6


def _build_body_section(doc: Document, sec: ParsedSection, cfg: ThesisConfig,
                         input_dir: Path, img_counter: dict,
                         parsed: ParsedDocument, fn_state: dict,
                         warnings: Optional[List[str]] = None) -> None:
    """构建正文章节"""
    # 标题
    if sec.level == 1:
        _add_heading1(doc, sec.title, cfg)
        # 新章节：递增章号，重置图片计数
        img_counter['chapter'] += 1
        img_counter['img'] = 0
        img_counter['table'] = 0
    elif sec.level == 2:
        _add_heading2(doc, sec.title, cfg)
    elif sec.level == 3:
        _add_heading3(doc, sec.title, cfg)

    # 正文内容
    for line in sec.content.split('\n'):
        _emit_content_line(doc, line, sec, cfg, input_dir, img_counter,
                           parsed, fn_state, warnings)


def _emit_content_line(doc: Document, line: str, sec: ParsedSection,
                       cfg: ThesisConfig, input_dir: Path, img_counter: dict,
                       parsed: ParsedDocument, fn_state: dict,
                       warnings: Optional[List[str]] = None) -> None:
    """输出一行内容：图片/图题/表格/公式/普通段落（正文与附录共用管线）"""
    trimmed = line.strip()
    if not trimmed or trimmed.startswith('#'):
        return

    # 清理 HTML 标签
    trimmed = re.sub(r'<br\s*/?>', '', trimmed, flags=re.IGNORECASE).strip()
    if not trimmed:
        return

    # 处理图片引用 ![caption](path)
    img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', trimmed)
    if img_match:
        caption = img_match.group(1)
        img_path = img_match.group(2)
        if not Path(img_path).is_absolute():
            img_path = input_dir / img_path
        img_counter['img'] += 1
        auto_cap = f"图{img_counter['chapter']}-{img_counter['img']}"
        if caption:
            auto_cap += f" {caption}"
        _add_image(doc, Path(img_path), auto_cap, cfg)
        return

    # 处理中文图片引用 "图 X-X 描述"（附录里可能是 "图A-1 描述"）
    if re.match(r'^图\s*[A-Za-z]?[-－]?\d', trimmed):
        img_dir = input_dir / "photo"
        if img_dir.exists():
            imgs = sorted(img_dir.iterdir(), key=lambda f: _extract_num(f.name))
            if img_counter['photo_idx'] < len(imgs):
                img_counter['img'] += 1
                auto_cap = f"图{img_counter['chapter']}-{img_counter['img']}"
                # 从原文提取描述（去掉已有的"图X-X"前缀）
                desc = re.sub(r'^图\s*[A-Za-z]?[-－]?[\d\-]+\s*', '', trimmed).strip()
                if desc:
                    auto_cap += f" {desc}"
                _add_image(doc, imgs[img_counter['photo_idx']], auto_cap, cfg)
                img_counter['photo_idx'] += 1
            else:
                _add_body_paragraph(doc, trimmed, cfg)
        else:
            _add_body_paragraph(doc, trimmed, cfg)
        return

    # 处理表格占位符 <!--TABLE:N-->
    if re.match(r'^<!--TABLE:(\d+)-->$', trimmed):
        table_idx = int(re.match(r'^<!--TABLE:(\d+)-->$', trimmed).group(1))
        if table_idx < len(sec.tables):
            td = sec.tables[table_idx]
            img_counter['table'] += 1
            # 如果没有表标题，自动生成 "表X-Y"
            if not td.caption:
                td.caption = f"表{img_counter['chapter']}-{img_counter['table']}"
            _build_table(doc, td, cfg)
        else:
            _add_body_paragraph(doc, f"[表格数据缺失: index={table_idx}]", cfg)

    # 处理公式占位符 <!--FORMULA:N-->（docx 导入，OMML 原样还原）
    elif _FORMULA_MARKER_RE.search(trimmed):
        _add_paragraph_with_formulas(doc, trimmed, cfg, parsed, fn_state,
                                     input_dir, warnings)

    else:
        # 普通段落：可解析的 [^N]/[N] 标记转为 Word 自动脚注
        _add_body_paragraph_with_notes(doc, trimmed, cfg, parsed, fn_state)


# 附录编号行：如「附录一」「附录2」「附录A」（规范：单独一行、左顶格）
_APPENDIX_NO_RE = re.compile(r'^附\s*录\s*[一二三四五六七八九十\dA-Za-z]{1,4}$')


def _build_appendix(doc: Document, sec: ParsedSection, cfg: ThesisConfig,
                    input_dir: Path, img_counter: dict,
                    parsed: ParsedDocument, fn_state: dict,
                    warnings: Optional[List[str]] = None) -> None:
    """构建附录（规范第12条）。

    - 标题「附  录」宋体三号加粗居中（ThesisBackTitle 样式，进目录）
    - 「附录N」编号单独一行左顶格；其下一行为该附录标题，居中、段前后1行
    - 其余内容与正文一致（表格/图片/公式/脚注管线全部可用）
    """
    _add_styled_heading(doc, "附  录", 'Thesis Back Title', cfg.spacing.line_spacing)

    lines = sec.content.split('\n')
    i = 0
    while i < len(lines):
        trimmed = lines[i].strip()
        if _APPENDIX_NO_RE.match(trimmed):
            # 编号行：左顶格、无首行缩进
            para = doc.add_paragraph()
            run = para.add_run(trimmed)
            _set_run_font(run, cfg.fonts.body_east_asian, cfg.fonts.body_latin,
                          cfg.fonts.body_size_pt, bold=True)
            _set_paragraph_spacing(para, line_spacing=cfg.spacing.line_spacing)
            # 下一非空行 = 附录标题：居中、段前1行、段后1行
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines):
                title = lines[j].strip()
                if (not _APPENDIX_NO_RE.match(title)
                        and not title.startswith(('|', '!', '<!--'))):
                    tpara = doc.add_paragraph()
                    tpara.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    trun = tpara.add_run(title)
                    _set_run_font(trun, cfg.fonts.body_east_asian,
                                  cfg.fonts.body_latin, cfg.fonts.body_size_pt,
                                  bold=True)
                    _set_paragraph_spacing(
                        tpara, before_pt=cfg.fonts.body_size_pt,
                        after_pt=cfg.fonts.body_size_pt,
                        line_spacing=cfg.spacing.line_spacing)
                    i = j + 1
                    continue
            i += 1
            continue
        _emit_content_line(doc, lines[i], sec, cfg, input_dir, img_counter,
                           parsed, fn_state, warnings)
        i += 1


def _extract_num(filename: str) -> int:
    """从文件名中提取数字用于排序"""
    m = re.search(r'\d+', filename)
    return int(m.group()) if m else 0


def _add_image(doc: Document, img_path: Path, caption: str, cfg: ThesisConfig,
                max_width_inches: float = 4) -> None:
    """插入图片和图标题"""
    if not img_path.exists():
        # 图片不存在时添加占位文本
        _add_body_paragraph(doc, f"[图片未找到: {img_path}]", cfg)
        return

    # 插入图片（居中）
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(img_path), width=Inches(max_width_inches))

    # 图标题（仿宋五号，居中）
    cap_para = doc.add_paragraph()
    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_para.paragraph_format.space_after = Pt(cfg.spacing.caption_after_pt)
    cap_run = cap_para.add_run(caption)
    _set_run_font(cap_run, cfg.fonts.caption_font, FONT_TIMES, cfg.fonts.caption_size_pt)


# 脚注标记：[^N]（md 脚注）或 [N]（句末引用参考文献的数字标）
_NOTE_MARKER_RE = re.compile(r'\[\^(\d+)\]|\[(\d{1,3})\]')

# 公式占位标记（docx 导入时抠出的 OMML，见 docx_import）
_FORMULA_MARKER_RE = re.compile(r'<!--FORMULA:(\d+)-->')


def _emit_text_with_notes(doc: Document, para, text: str, cfg: ThesisConfig,
                          parsed: ParsedDocument, fn_state: dict) -> None:
    """向已有段落写入文本，可解析的脚注标记转为 Word 自动脚注（规范第10条）。

    - [^N]：脚注内容取 md 中的 [^N]: 定义
    - [N]：脚注内容取参考文献第 N 条（格式与参考文献一致）
    解析不了的标记原样保留。
    """
    refs = parsed.get_references()

    def _resolve(m: re.Match) -> Optional[str]:
        if m.group(1):
            return parsed.footnotes.get(int(m.group(1)))
        n = int(m.group(2))
        if 1 <= n <= len(refs):
            return re.sub(r'^\[\d+\]\s*', '', refs[n - 1]).strip()
        return None

    usable = [(m, _resolve(m)) for m in _NOTE_MARKER_RE.finditer(text)]
    usable = [(m, t) for m, t in usable if t]

    pos = 0
    for m, fn_text in usable:
        if m.start() > pos:
            run = para.add_run(text[pos:m.start()])
            _set_run_font(run, cfg.fonts.body_east_asian,
                          cfg.fonts.body_latin, cfg.fonts.body_size_pt)
        fn_state['next_id'] = add_footnotes(doc, para, [fn_text], fn_state['next_id'])
        pos = m.end()
    if pos < len(text):
        run = para.add_run(text[pos:])
        _set_run_font(run, cfg.fonts.body_east_asian,
                      cfg.fonts.body_latin, cfg.fonts.body_size_pt)


def _add_body_paragraph_with_notes(doc: Document, text: str, cfg: ThesisConfig,
                                   parsed: ParsedDocument, fn_state: dict) -> None:
    """添加正文段落（脚注标记自动转 Word 脚注）"""
    if not _NOTE_MARKER_RE.search(text):
        _add_body_paragraph(doc, text, cfg)
        return
    para = doc.add_paragraph()
    _emit_text_with_notes(doc, para, text, cfg, parsed, fn_state)
    _set_paragraph_spacing(
        para, line_spacing=cfg.spacing.line_spacing,
        first_line_chars=cfg.spacing.first_line_indent_chars,
    )


def _append_formula(para, n: int, input_dir: Path,
                    warnings: Optional[List[str]]) -> bool:
    """把 formula/N.xml 的 OMML 原样挂进段落，失败返回 False 并记 warning"""
    xml_path = input_dir / "formula" / f"{n}.xml"
    try:
        el = etree.fromstring(xml_path.read_bytes())
        para._p.append(el)
        return True
    except (OSError, etree.XMLSyntaxError) as exc:
        if warnings is not None:
            warnings.append(f"第 {n} 处公式还原失败（{exc.__class__.__name__}），已跳过")
        return False


def _add_paragraph_with_formulas(doc: Document, text: str, cfg: ThesisConfig,
                                 parsed: ParsedDocument, fn_state: dict,
                                 input_dir: Path,
                                 warnings: Optional[List[str]]) -> None:
    """含 <!--FORMULA:N--> 标记的段落：文本与 OMML 公式交替写入。

    整段只有公式标记时按独立公式排（居中、无首行缩进）。
    """
    standalone = not _FORMULA_MARKER_RE.sub('', text).strip()
    para = doc.add_paragraph()

    pos = 0
    for m in _FORMULA_MARKER_RE.finditer(text):
        if m.start() > pos:
            _emit_text_with_notes(doc, para, text[pos:m.start()], cfg, parsed, fn_state)
        _append_formula(para, int(m.group(1)), input_dir, warnings)
        pos = m.end()
    if pos < len(text):
        _emit_text_with_notes(doc, para, text[pos:], cfg, parsed, fn_state)

    if standalone:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_paragraph_spacing(para, line_spacing=cfg.spacing.line_spacing)
    else:
        _set_paragraph_spacing(
            para, line_spacing=cfg.spacing.line_spacing,
            first_line_chars=cfg.spacing.first_line_indent_chars,
        )


def _set_cell_border(cell, **kwargs) -> None:
    """设置单元格边框。

    kwargs: top, bottom, left, right — 每个值为 dict，包含:
        sz (int): 线宽（1/8 磅），如 12 = 1.5pt
        val (str): 线型，默认 'single'
        color (str): 颜色，默认 '000000'
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
        tcPr.append(tcBorders)
    for edge, attrs in kwargs.items():
        el = tcBorders.find(qn(f'w:{edge}'))
        if el is None:
            el = parse_xml(f'<w:{edge} {nsdecls("w")}/>')
            tcBorders.append(el)
        el.set(qn('w:val'), attrs.get('val', 'single'))
        el.set(qn('w:sz'), str(attrs.get('sz', 4)))
        el.set(qn('w:color'), attrs.get('color', '000000'))
        el.set(qn('w:space'), '0')


def _build_table(doc: Document, table_data: TableData, cfg: ThesisConfig) -> None:
    """构建三线表格。

    三线表格规则：
    - 顶线：粗线（表头上方）
    - 栏目线：粗线（表头下方）
    - 底线：粗线（表格最下方）
    - 无竖线，无其他横线
    - 表标题在表格上方：仿宋 五号(10.5pt) 居中
    """
    headers = table_data.headers
    rows = table_data.rows
    if not headers:
        return

    num_cols = len(headers)

    # ① 表标题（在表格上方）
    if table_data.caption:
        cap_para = doc.add_paragraph()
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_para.paragraph_format.space_before = Pt(6)
        cap_para.paragraph_format.space_after = Pt(3)
        cap_run = cap_para.add_run(table_data.caption)
        _set_run_font(cap_run, FONT_FANGSONG, FONT_TIMES, 10.5)  # 五号

    # ② 创建表格
    num_row_data = len(rows) if rows else 0
    table = doc.add_table(rows=1 + num_row_data, cols=num_cols)

    # ③ 表格宽度设为页面可用宽度的 100%
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    # 移除已有 tblW
    for old in tblPr.findall(qn('w:tblW')):
        tblPr.remove(old)
    tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:type="pct" w:w="5000"/>')
    tblPr.append(tblW)
    # 自动适应窗口
    for old in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(old)
    tblLayout = parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="autofit"/>')
    tblPr.append(tblLayout)

    # ④ 清除表格默认边框（设为无边框）
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    no_border_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="none" w:sz="0" w:color="auto" w:space="0"/>'
        '<w:left w:val="none" w:sz="0" w:color="auto" w:space="0"/>'
        '<w:bottom w:val="none" w:sz="0" w:color="auto" w:space="0"/>'
        '<w:right w:val="none" w:sz="0" w:color="auto" w:space="0"/>'
        '<w:insideH w:val="none" w:sz="0" w:color="auto" w:space="0"/>'
        '<w:insideV w:val="none" w:sz="0" w:color="auto" w:space="0"/>'
        '</w:tblBorders>'
    )
    tblPr.append(parse_xml(no_border_xml))

    # ⑤ 填写表头行
    header_row = table.rows[0]
    thick = {'sz': 12, 'val': 'single', 'color': '000000'}  # 12 × 1/8pt = 1.5pt
    none_b = {'sz': 0, 'val': 'none', 'color': 'auto'}
    for j, htext in enumerate(headers):
        cell = header_row.cells[j]
        # 清空默认段落
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(htext)
        _set_run_font(run, FONT_SIMSUN, FONT_TIMES, 10.5)  # 五号
        # 表头：顶线(粗) + 底线(粗)，左右无
        _set_cell_border(cell,
                         top=thick, bottom=thick,
                         left=none_b, right=none_b)

    # ⑥ 填写数据行
    for r_idx, row_data in enumerate(rows):
        row = table.rows[1 + r_idx]
        is_last = (r_idx == len(rows) - 1)
        for j in range(num_cols):
            cell_text = row_data[j] if j < len(row_data) else ""
            cell = row.cells[j]
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(cell_text)
            _set_run_font(run, FONT_SIMSUN, FONT_TIMES, 10.5)  # 五号
            # 数据行：默认无边框
            border_kwargs = {'top': none_b, 'bottom': none_b,
                             'left': none_b, 'right': none_b}
            # 最后一行加底线(粗)
            if is_last:
                border_kwargs['bottom'] = thick
            _set_cell_border(cell, **border_kwargs)

    # ⑦ 表格后添加空段落（间距）
    after_para = doc.add_paragraph()
    after_para.paragraph_format.space_before = Pt(3)
    after_para.paragraph_format.space_after = Pt(6)


# 需要导入 Inches
from docx.shared import Inches


# ═══════════════════════════════════════════
# 主转换函数
# ═══════════════════════════════════════════

def _deep_merge(base: Optional[dict], override: Optional[dict]) -> Optional[dict]:
    """按嵌套键合并两个配置字典（override 优先），任一为空时返回另一个"""
    if not base:
        return override
    if not override:
        return base
    out = dict(base)
    for k, v in override.items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = _deep_merge(out[k], v)
        else:
            out[k] = v
    return out


def _resource_dir_for_input(input_md: Path) -> Path:
    """Return the folder that holds the input's referenced assets."""
    return input_md.parent


def convert(
    input_md: str | Path,
    template_path: str | Path,
    config: Optional[dict] = None,
    output_path: Optional[str | Path] = None,
    output_dir: Optional[str | Path] = None,
    fast: bool = False,
) -> dict:
    """将 Markdown 文件转换为符合山西财经大学格式规范的 DOCX。

    Args:
        input_md: 输入 Markdown 文件路径
        template_path: 学校模板 DOCX 路径
        config: 可选的格式配置覆盖
        output_path: 输出 DOCX 路径（默认自动生成）
        output_dir: output_path 未指定时的输出目录（默认 output/；预览用隔离目录）
        fast: 跳过 Word 分页后处理（奇偶页空白填充与目录域落盘），全程零
            Word 启动。仅供预览：预览 PDF 由 docx2pdf.ps1 -UpdateFields
            在导出会话内更新域，目录页码仍正确，只是没有奇偶对齐空白页；
            产出的 .docx 目录域是未展开的占位，不可交付。

    Returns:
        包含 success, output_path, pages, warnings 的字典
        （pages：完整路径为 Word 实测值，fast 路径为段落数估算）
    """
    input_md = Path(input_md)
    template_path = Path(template_path)
    warnings: List[str] = []

    # 解析配置：模板档案里的 config 是该模板（该校规范）的格式默认值，
    # 请求传入的 config 在其之上覆盖。这样直调 API / 命令行也能拿到正确格式，
    # 不依赖前端先把档案合并进面板。
    profile_config = (load_template_profile(template_path) or {}).get('config')
    merged_config = _deep_merge(profile_config, config)
    cfg = ThesisConfig.from_dict(merged_config) if merged_config else ThesisConfig()

    # 解析 Markdown
    parsed = parse_markdown(input_md)
    # Relative image links are resolved from the input Markdown's own folder.
    input_dir = _resource_dir_for_input(input_md)

    # 确定输出路径
    if output_path is None:
        safe_title = re.sub(r'[<>:"/\\|?*]', '_', parsed.title)[:80]
        output_path = Path(output_dir or "output") / f"{safe_title}.docx"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # 创建文档
    doc = Document()

    # 设置页面
    _setup_page(doc, cfg)

    # 设置文档默认值、默认字体和标题/目录样式（与学校模板一致）
    _setup_doc_defaults(doc, cfg)
    _setup_default_font(doc, cfg)
    _ensure_heading_styles(doc, cfg)

    # ── 构建各章节，同时跟踪每个分节的类型和奇偶要求 ──
    # section_types / section_parities 与 doc.sections 一一对应（合并模板后封面分节会前插）
    # 分节符一律用「下一页」；奇偶起始页由保存后的 Word 分页后处理补真实空白页实现
    # （规范：空白填充页也要按规则显示页眉页码，Word 的 odd/even 分节自动空白页做不到）
    #
    # ⚠️ oddPage 方案已三次否决，勿再尝试（历史：1572fbf 全 oddPage → f5cd354 回滚）。
    # 2026-07-20 受控实验定论：取真实产物删 8 个填充段落、sect1-6 改 oddPage/sect8 改
    # evenPage 后由 Word 导 PDF——页数与各节落点和基准完全一致（18 页），但 4 张自动
    # 空白页（p6/p8/p12/p14）全部变成零绘图对象的裸白页：无页眉、无页码（基准上是
    # 页眉+II/IV/2/4）。规范第 5 行要求空白页按规定带页眉页码 → 此路在规范层面封死。
    # 实验材料与逐页比对脚本见 new_sxupaper/HANDOFF.md §4 实验 A。
    section_types: List[str] = []
    section_parities: List[str] = []

    # 如果有封面模板，在最前面预留一个空分节给封面内容
    has_cover = cfg.sections.cover and template_path.exists()
    if has_cover:
        add_section_break(doc, "new_page")   # section 0 → 留给封面

    # ── 前置部分：中文摘要 / 英文摘要 / 目录 ──
    if cfg.sections.abstract_cn:
        section_types.append("front_matter")
        section_parities.append("odd")       # 中文摘要首页奇数页
        _build_abstract_cn(doc, parsed, cfg)
        add_section_break(doc, "new_page")

    if cfg.sections.abstract_en:
        section_types.append("front_matter")
        section_parities.append("odd")       # 英文摘要首页奇数页
        _build_abstract_en(doc, parsed, cfg)
        add_section_break(doc, "new_page")

    if cfg.sections.toc:
        section_types.append("front_matter")
        section_parities.append("odd")       # 目录首页奇数页
        _build_toc(doc, cfg)
        add_section_break(doc, "new_page")

    # ── 正文 ──
    img_counter = {'chapter': 0, 'img': 0, 'photo_idx': 0, 'table': 0}
    fn_state = {'next_id': 1}   # Word 脚注 ID 全局递增（显示编号由每页重编控制）
    if cfg.sections.body:
        section_types.append("body")
        section_parities.append("odd")       # 正文首页奇数页
        for sec in parsed.get_body_sections():
            _build_body_section(doc, sec, cfg, input_dir, img_counter, parsed,
                                fn_state, warnings)

    # ── 参考文献（独立分节）──
    if cfg.sections.references:
        refs = parsed.get_references()
        if refs:
            add_section_break(doc, "new_page")
            section_types.append("body")
            section_parities.append("odd")   # 参考文献首页奇数页
            _add_heading1(doc, "参考文献", cfg)
            _build_references(doc, refs, cfg)
        else:
            warnings.append("未找到参考文献")

    # ── 附录（独立分节，内容与正文同管线：表格/图片/公式/脚注可用）──
    if cfg.sections.appendix and parsed.has_appendix():
        add_section_break(doc, "new_page")
        section_types.append("body")
        section_parities.append("odd")       # 附录首页奇数页
        sec = parsed.find_section("附录")
        if sec:
            _build_appendix(doc, sec, cfg, input_dir, img_counter,
                            parsed, fn_state, warnings)
    elif cfg.sections.appendix:
        warnings.append(
            "已勾选附录但文档中没有附录内容，已跳过附录页"
            "（md 请添加「# 附录」章节；Word 请添加「附 录」标题段落）")

    # ── 致谢（独立分节）──
    if cfg.sections.acknowledgment:
        add_section_break(doc, "new_page")
        section_types.append("body")
        section_parities.append("odd")       # 致谢首页奇数页
        ack = parsed.get_acknowledgment_content()
        _add_styled_heading(doc, "致  谢", 'Thesis Back Title', cfg.spacing.line_spacing)
        if ack:
            # 逐段输出（整块塞一个段落会把多段内容压成一段）
            for ack_line in ack.split('\n'):
                if ack_line.strip():
                    _add_body_paragraph(doc, ack_line.strip(), cfg)
        else:
            _add_body_paragraph(
                doc,
                "本论文的写作过程是一次系统学习学术研究方法的宝贵经历。"
                "在导师的悉心指导下，我对相关领域有了更深入的认识，"
                "也初步掌握了文献研究、案例比较等基本方法。"
                "感谢导师在选题确定、资料搜集、论文修改各环节给予的耐心指导，"
                "感谢同学们在讨论交流中提供的启发与帮助。"
                "今后将继续努力，不断提升自身的学术素养与研究能力。",
                cfg,
            )

    # 如果有封面模板，在末尾预留两个空分节：
    # 先是附件2/成绩评定表，最后是校徽封底页。
    # 封底落偶数页＝最后一张纸的背面，双面打印装订后合上正好朝外露出
    if has_cover:
        add_section_break(doc, "new_page")
        section_types.append("cover")    # 附件2 / 成绩评定表
        section_parities.append("any")
        add_section_break(doc, "new_page")
        section_types.append("cover")    # 校徽封底页
        section_parities.append("even")

    # ── 合并模板封面/封底 ──
    if has_cover:
        try:
            # 年级（学号前4位=入学年）随元数据提供给封面填空，
            # 供「年级」栏标签映射（南大等封面 年级/学号 同行双栏）
            merge_metadata = dict(parsed.metadata)
            merge_metadata.setdefault(
                "grade", parsed.get_meta("studentId", "202300000000")[:4])
            merge_template(doc, template_path, merge_metadata, warnings)
            # 模板封面被插入到最前面的空分节（section 0），类型为 cover
            section_types = ["cover"] + section_types
            section_parities = ["any"] + section_parities
        except Exception as e:
            warnings.append(f"模板合并失败: {e}")

    # ── 设置页眉页脚（传入分节类型列表）──
    grade = parsed.get_meta("studentId", "202300000000")[:4]
    setup_headers_footers(doc, cfg, grade, section_types,
                          title=parsed.get_meta("title", ""))

    # 批注已禁用（预览和打印时不需要批注）

    # 保存
    doc.save(str(output_path))

    # ── Word 分页后处理：按规范把各章节首页调整到奇/偶数页，
    #    插入的是上一分节内的真实空白页（带该分节的页眉页码）──
    #    fast（预览快速通道）跳过：奇偶空白页是成稿要求，预览不需要
    real_pages: Optional[int] = None
    if not fast:
        real_pages = finalize_docx(output_path, section_parities, warnings)

    return {
        "success": True,
        "output_path": str(output_path),
        "preview_url": f"/api/preview/{output_path.stem}",
        "pages": real_pages if real_pages else _estimate_pages(doc),
        "warnings": warnings,
        # A fast preview is the same pre-pagination DOCX as a final export.
        # The API may cache it and later finalize a copy, avoiding a second
        # complete Markdown -> DOCX build after the user has already previewed.
        "section_parities": section_parities,
    }


_SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"


def _run_word_script(script_name: str, args: List[str],
                     timeout: int = 180) -> Optional[str]:
    """运行 scripts/ 下的 Word COM 脚本，成功返回 stdout，失败返回 None"""
    script = _SCRIPTS_DIR / script_name
    if not script.exists():
        return None
    result = subprocess.run(
        ["powershell.exe", "-ExecutionPolicy", "Bypass", "-File", str(script)] + args,
        capture_output=True, timeout=timeout,
        **hidden_process_kwargs(),
    )
    if result.returncode != 0:
        return None
    return result.stdout.decode("gbk", errors="replace")


def _word_pass(docx_path: Path, *, update_fields: bool = False,
               update_page_numbers: bool = False,
               measure: bool = False) -> Optional[str]:
    """一次 Word 会话组合执行 更新域 / 刷目录页码 / 测量分节。

    每次 Word 冷启动约 8s，把原来分两次启动的「更新+测量」并进一个
    会话是分页后处理最大的提速来源。失败返回 None。
    """
    # The old persistent-session attempt edited page breaks inside Word and was
    # correctly reverted for unstable pagination. This worker is different: all
    # structural edits remain Python/XML writes between passes; Word only opens,
    # updates/measures, saves and closes one document. On any fault retain the
    # one-shot PowerShell path below.
    try:
        from ..api.word_preview_worker import WORD_PREVIEW_WORKER
        return WORD_PREVIEW_WORKER.pass_docx(
            docx_path, update_fields=update_fields,
            update_page_numbers=update_page_numbers, measure=measure)
    except Exception as exc:
        print(f"[word-pass-worker] fallback to one-shot Word: {exc}")

    args = ["-Docx", str(docx_path)]
    if update_fields:
        args.append("-UpdateFields")
    if update_page_numbers:
        args.append("-UpdatePageNumbers")
    if measure:
        args.append("-Measure")
    return _run_word_script("word_pass.ps1", args)


def _parse_measure(out: str) -> tuple[Optional[Dict[int, int]], Optional[int]]:
    """解析 word_pass -Measure 输出 → (各分节起始物理页, 实测总页数)"""
    pages: Dict[int, int] = {}
    total: Optional[int] = None
    for m in re.finditer(r'^(\d+):(\d+)', out, re.MULTILINE):
        pages[int(m.group(1))] = int(m.group(2))
    m = re.search(r'^PAGES:(\d+)', out, re.MULTILINE)
    if m:
        total = int(m.group(1))
    return (pages or None), total


def _insert_filler_breaks(docx_path: Path, section_indices: List[int]) -> None:
    """在指定分节（1 起）的「上一分节」末尾插入显式分页符段落。

    XML 层操作，位置确定：插到上一分节的 sectPr 段落之前，空白页留在
    上一分节内，因此保有其页眉页脚（规范要求填充页也显示页眉页码）。
    """
    doc = Document(str(docx_path))
    body = doc.element.body
    # 收集各分节的结束 sectPr 段落（第 k 个段落级 sectPr 结束第 k+1 个分节，0起）
    sect_paras = [el for el in body
                  if el.tag == qn('w:p') and el.find(qn('w:pPr')) is not None
                  and el.find(qn('w:pPr')).find(qn('w:sectPr')) is not None]
    for sec_no in section_indices:
        prev_idx = sec_no - 2   # 上一分节的 sectPr 段落下标（0起）
        if not (0 <= prev_idx < len(sect_paras)):
            continue
        brk = parse_xml(
            f'<w:p {nsdecls("w")}><w:r><w:br w:type="page"/></w:r></w:p>')
        sect_paras[prev_idx].addprevious(brk)
    doc.save(str(docx_path))


def finalize_docx(docx_path: Path, section_parities: List[str],
                  warnings: List[str], *, fields_already_updated: bool = False) -> Optional[int]:
    """Turn a fast/pre-pagination DOCX into a deliverable DOCX in place.

    This is deliberately a narrow public wrapper around the established Word
    pagination pass.  It lets callers reuse an immutable cached source DOCX;
    the cached file is always copied first, because this routine inserts real
    filler-page breaks and updates the TOC cache on disk.
    """
    return _fix_pagination(docx_path, section_parities, warnings,
                           fields_already_updated=fields_already_updated)


def _fix_pagination(docx_path: Path, parities: List[str], warnings: List[str],
                    *, fields_already_updated: bool = False) -> Optional[int]:
    """按分节奇偶要求补真实空白页（带上一分节的页眉页码）。

    流程：「Word 会话（更新域/刷页码 + 测量）→ python-docx 在 XML 层
    插分页符」循环至收敛。比在 Word COM 里热编辑插分页符可靠
    （热编辑时插入位置和分页结果都不稳定）。
    parities 与 doc.sections 一一对应，取值 odd / even / any。

    Word 启动次数 = 1 + 插页轮数（无需插页 1 次、典型 2 次、最坏 4 次）：
    首轮会话做全量域更新+测量；每次插页后的会话「刷目录页码+测量」合并，
    因此收敛那一轮的目录页码已经刷新过，不再补启动。旧实现每步单起
    一次 Word，同文档要 2（无插页）～6（最坏）次。

    Returns:
        Word 实测的总页数；测量全部失败时 None。
    """
    docx_path = Path(docx_path).resolve()
    total_pages: Optional[int] = None
    try:
        # A saved fast-preview cache has already performed the full field pass.
        # It is safe to skip only that duplicated work here: this first pass
        # still repaginates and measures, and every later filler-page round
        # continues to update TOC page numbers and measure as before.
        out = _word_pass(docx_path, update_fields=not fields_already_updated,
                         measure=True)
        if out is None:
            warnings.append("奇偶页调整失败：Word 更新目录域未成功")
            return None

        def _plan(pages: Dict[int, int]) -> List[int]:
            """按测量结果规划需要补空白页的分节。

            关键：同一轮内上游插入的空白页会把下游分节整体再推一页，
            必须按累计位移算「有效页码」再判断奇偶，否则上游修正与
            下游自身修正互相抵消，循环永不收敛。
            """
            inserts = []
            shift = 0
            for sec_no, target in enumerate(parities, 1):
                page = pages.get(sec_no)
                if page is None:
                    continue
                effective = page + shift
                if target in ("odd", "even") and (target == "odd") != (effective % 2 == 1):
                    inserts.append(sec_no)
                    shift += 1
            return inserts

        for _ in range(3):
            pages, total = _parse_measure(out)
            if total is not None:
                total_pages = total
            if pages is None:
                warnings.append("奇偶页调整失败：Word 测量分节页码未成功")
                return total_pages
            plan = _plan(pages)
            if not plan:
                break
            _insert_filler_breaks(docx_path, plan)
            # 插页移动了显示页码：刷新目录页码，并在同一会话里复测
            out = _word_pass(docx_path, update_page_numbers=True, measure=True)
            if out is None:
                warnings.append("奇偶页调整失败：Word 测量分节页码未成功")
                return total_pages
        else:
            # 循环耗尽：末轮测量已在手（插页后复测过），如实报告残留违规
            pages, total = _parse_measure(out)
            if total is not None:
                total_pages = total
            for sec_no in (_plan(pages) if pages else []):
                target_cn = "奇数" if parities[sec_no - 1] == "odd" else "偶数"
                warnings.append(
                    f"第 {sec_no} 分节未能调整到{target_cn}页起始，请打印前检查")
        return total_pages
    except (subprocess.TimeoutExpired, FileNotFoundError) as e:
        warnings.append(f"奇偶页调整跳过: {e}")
        return total_pages


def _setup_page(doc: Document, cfg: ThesisConfig) -> None:
    """设置页面尺寸、页边距和文档网格"""
    section = doc.sections[0]
    section.page_width = Cm(cfg.page.width_cm)
    section.page_height = Cm(cfg.page.height_cm)
    section.top_margin = Cm(cfg.page.margin_top_cm)
    section.bottom_margin = Cm(cfg.page.margin_bottom_cm)
    section.left_margin = Cm(cfg.page.margin_left_cm)
    section.right_margin = Cm(cfg.page.margin_right_cm)
    section.header_distance = Cm(1.27)  # 页眉距边界
    section.footer_distance = Cm(1.27)  # 页脚距边界

    # 文档网格：与学校标准文档一致，linePitch=1 即不启用行网格。
    # （若启用 44字×43行 网格，1.5倍行距的行会被吸附到 2 格 31.2pt，导致封面表格膨胀跨页）
    sectPr = section._sectPr
    for old in sectPr.findall(qn('w:docGrid')):
        sectPr.remove(old)
    docGrid = etree.SubElement(sectPr, qn('w:docGrid'))
    docGrid.set(qn('w:linePitch'), '1')

    # 封面分节从奇数页开始（与标准文档一致）
    type_elem = sectPr.find(qn('w:type'))
    if type_elem is None:
        type_elem = parse_xml(f'<w:type {nsdecls("w")} w:val="oddPage"/>')
        sectPr.insert(0, type_elem)
    else:
        type_elem.set(qn('w:val'), 'oddPage')


def _setup_doc_defaults(doc: Document, cfg: ThesisConfig) -> None:
    """重写 docDefaults，与学校模板一致。

    python-docx 自带模板的默认值是 Calibri 11pt、段后 10pt、1.15 倍行距，
    模板拷贝过来的段落（未显式设置间距的空行等）会继承这些值被撑高，
    导致封面/说明页内容溢出。学校模板为 宋体/Times 小四、无段后、1.25 倍行距。
    """
    styles_el = doc.styles.element
    old_dd = styles_el.find(qn('w:docDefaults'))
    new_dd = parse_xml(
        f'<w:docDefaults {nsdecls("w")}>'
        '<w:rPrDefault><w:rPr>'
        f'<w:rFonts w:ascii="{cfg.fonts.body_latin}" w:eastAsia="宋体"'
        f' w:hAnsi="{cfg.fonts.body_latin}" w:cs="{cfg.fonts.body_latin}"/>'
        f'<w:sz w:val="{int(cfg.fonts.body_size_pt * 2)}"/>'
        f'<w:szCs w:val="{int(cfg.fonts.body_size_pt * 2)}"/>'
        '<w:lang w:val="en-US" w:eastAsia="zh-CN" w:bidi="ar-SA"/>'
        '</w:rPr></w:rPrDefault>'
        '<w:pPrDefault><w:pPr>'
        '<w:spacing w:line="300" w:lineRule="auto"/>'
        '</w:pPr></w:pPrDefault>'
        '</w:docDefaults>'
    )
    if old_dd is not None:
        styles_el.replace(old_dd, new_dd)
    else:
        styles_el.insert(0, new_dd)


def _replace_style(doc: Document, style_id: str, xml: str) -> None:
    """删除同 styleId 的旧样式并追加新定义"""
    styles_el = doc.styles.element
    for st in styles_el.findall(qn('w:style')):
        if st.get(qn('w:styleId')) == style_id:
            styles_el.remove(st)
    styles_el.append(parse_xml(xml))


def _ensure_heading_styles(doc: Document, cfg: ThesisConfig) -> None:
    """定义标题和目录段落样式。

    标题必须走样式而非 run 直接格式：Word 更新目录域时会把标题的
    「直接字符格式」（如加粗）带进目录条目，样式格式则不会——
    规范要求目录正文为宋体小四不加粗。
    """
    f = cfg.fonts
    sp = cfg.spacing

    def _sz(pt: float) -> int:
        return int(pt * 2)

    def _tw(pt: float) -> int:
        return int(pt * 20)

    h1_bold = '<w:b/>' if f.heading1_bold else ''
    _replace_style(doc, 'Heading1', (
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="Heading1">'
        '<w:name w:val="heading 1"/><w:basedOn w:val="Normal"/>'
        '<w:next w:val="Normal"/><w:qFormat/>'
        '<w:pPr><w:keepNext/>'
        f'<w:spacing w:before="{_tw(sp.heading1_before_pt)}" w:after="{_tw(sp.heading1_after_pt)}"/>'
        '<w:jc w:val="center"/><w:outlineLvl w:val="0"/></w:pPr>'
        f'<w:rPr><w:rFonts w:ascii="{f.body_latin}" w:hAnsi="{f.body_latin}"'
        f' w:eastAsia="{f.heading1_font}"/>{h1_bold}'
        f'<w:sz w:val="{_sz(f.heading1_size_pt)}"/><w:szCs w:val="{_sz(f.heading1_size_pt)}"/></w:rPr>'
        '</w:style>'
    ))

    h2_bold = '<w:b/>' if f.heading2_bold else ''
    _replace_style(doc, 'Heading2', (
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="Heading2">'
        '<w:name w:val="heading 2"/><w:basedOn w:val="Normal"/>'
        '<w:next w:val="Normal"/><w:qFormat/>'
        '<w:pPr><w:keepNext/>'
        f'<w:spacing w:before="{_tw(sp.heading2_before_pt)}" w:after="{_tw(sp.heading2_after_pt)}"/>'
        '<w:outlineLvl w:val="1"/></w:pPr>'
        f'<w:rPr><w:rFonts w:ascii="{f.body_latin}" w:hAnsi="{f.body_latin}"'
        f' w:eastAsia="{f.heading2_font}"/>{h2_bold}'
        f'<w:sz w:val="{_sz(f.heading2_size_pt)}"/><w:szCs w:val="{_sz(f.heading2_size_pt)}"/></w:rPr>'
        '</w:style>'
    ))

    h3_bold = '<w:b/>' if f.heading3_bold else ''
    _replace_style(doc, 'Heading3', (
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="Heading3">'
        '<w:name w:val="heading 3"/><w:basedOn w:val="Normal"/>'
        '<w:next w:val="Normal"/><w:qFormat/>'
        '<w:pPr><w:keepNext/>'
        f'<w:spacing w:before="{_tw(sp.heading3_before_pt)}" w:after="{_tw(sp.heading3_after_pt)}"/>'
        '<w:outlineLvl w:val="2"/></w:pPr>'
        f'<w:rPr><w:rFonts w:ascii="{f.body_latin}" w:hAnsi="{f.body_latin}"'
        f' w:eastAsia="{f.heading3_font}"/>{h3_bold}'
        f'<w:sz w:val="{_sz(f.heading3_size_pt)}"/><w:szCs w:val="{_sz(f.heading3_size_pt)}"/></w:rPr>'
        '</w:style>'
    ))

    # 致谢/附录标题样式（宋体三号加粗居中，进目录一级）
    _replace_style(doc, 'ThesisBackTitle', (
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="ThesisBackTitle">'
        '<w:name w:val="Thesis Back Title"/><w:basedOn w:val="Normal"/>'
        '<w:next w:val="Normal"/><w:qFormat/>'
        '<w:pPr><w:keepNext/>'
        f'<w:spacing w:before="{_tw(sp.heading1_before_pt)}" w:after="{_tw(sp.heading1_after_pt)}"/>'
        '<w:jc w:val="center"/><w:outlineLvl w:val="0"/></w:pPr>'
        f'<w:rPr><w:rFonts w:ascii="{f.body_latin}" w:hAnsi="{f.body_latin}" w:eastAsia="{FONT_SIMSUN}"/>'
        '<w:b/><w:sz w:val="32"/><w:szCs w:val="32"/></w:rPr>'
        '</w:style>'
    ))

    # 目录条目样式（规范：目录正文宋体小四；一级顶格，二级左缩进2字符）
    _replace_style(doc, 'TOC1', (
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="TOC1">'
        '<w:name w:val="toc 1"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/>'
        '<w:uiPriority w:val="39"/>'
        f'<w:rPr><w:rFonts w:ascii="{f.body_latin}" w:hAnsi="{f.body_latin}" w:eastAsia="{FONT_SIMSUN}"/>'
        f'<w:sz w:val="{_sz(f.body_size_pt)}"/></w:rPr>'
        '</w:style>'
    ))
    _replace_style(doc, 'TOC2', (
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="TOC2">'
        '<w:name w:val="toc 2"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/>'
        '<w:uiPriority w:val="39"/>'
        '<w:pPr><w:ind w:leftChars="200" w:left="480"/></w:pPr>'
        f'<w:rPr><w:rFonts w:ascii="{f.body_latin}" w:hAnsi="{f.body_latin}" w:eastAsia="{FONT_SIMSUN}"/>'
        f'<w:sz w:val="{_sz(f.body_size_pt)}"/></w:rPr>'
        '</w:style>'
    ))


def _setup_default_font(doc: Document, cfg: ThesisConfig) -> None:
    """设置文档默认字体"""
    style = doc.styles['Normal']
    font = style.font
    font.name = cfg.fonts.body_latin
    font.size = Pt(cfg.fonts.body_size_pt)
    # 设置东亚字体
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rpr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), cfg.fonts.body_east_asian)


def _estimate_pages(doc: Document) -> int:
    """粗略估算页数（基于段落数量）"""
    para_count = len(doc.paragraphs)
    # A4 约 30 行/页，每行约 35 字
    # 粗略按每页 25 个段落估算
    return max(1, para_count // 25)
