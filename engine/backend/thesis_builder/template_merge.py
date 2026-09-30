"""模板合并：从学校模板 DOCX 提取封面+封底，合并到输出文档。

对应 C# 中的 MergeTemplate / FillMajorCode / FillCoverTable / CopyImages 等函数。
"""

from __future__ import annotations

import io
import json
import re
from copy import deepcopy
from pathlib import Path
from typing import Dict, List, Optional

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part


# ═══════════════════════════════════════════
# 模板格式档案（template/<模板名>.json）
# ═══════════════════════════════════════════
# 每份模板可带一份同名 JSON 档案描述它的锚点与填充规则，
# 没有档案时使用山财默认值（向后兼容）。档案结构：
# {
#   "name": "显示名",
#   "merge": {
#     "cover_start_anchor": "学校代码" | null,   # null=从文档开头
#     "explanation_anchor": "说明" | null,       # null=此模板无说明页
#     "front_end_anchors": ["摘要"],             # 依次找，全无则退回首个段级分节符
#     "tail": true,                              # false=此模板无尾部（评定表/封底）
#     "tail_form_anchors": ["附件", "成绩评定表"],
#     "major_code_label": "专业代码" | null,
#     "cover_labels": {"标签": "元数据键"},       # 增量合并进内置映射（表格填充）
#     "paragraph_labels": {"标签": "元数据键"},    # 段落式填空（「学生姓名：___」）
#     "placeholders": [{"pattern": "^X{6,}$", "value": "title"}]  # 正则占位符替换
#   },
#   "config": { ... }   # 选中该模板时前端自动带出的格式默认值（可省略）
# }

_DEFAULT_COVER_LABELS = {
    '中文题目': 'title',
    '论文题目': 'title',    # 旧版模板（附件4）等用此标签
    '课题名称': 'title',
    '英文题目': 'engTitle',
    '姓名': 'name',
    '学生姓名': 'name',
    '学号': 'studentId',
    '班级': 'class',
    '专业': 'major',
    '学院': 'college',
    '指导教师': 'advisor',
    '完成时间': 'date',
    '起止日期': 'date',
    '日期': 'date',
}

_DEFAULT_MERGE_OPTS = {
    'cover_start_anchor': '学校代码',
    'explanation_anchor': '说明',
    'front_end_anchors': ['摘要'],
    'tail': True,
    'tail_form_anchors': ['附件', '成绩评定表'],
    'major_code_label': '专业代码',
    'cover_labels': {},
    'paragraph_labels': {},
    'placeholders': [],
    'inherit_page_setup': 'frame',   # frame | full | none，见 merge_template 内注释
    'title_layout': {},              # 标题栏排版参数，见 _DEFAULT_TITLE_LAYOUT
}


def load_template_profile(template_path: str | Path,
                          warnings: Optional[List[str]] = None) -> dict:
    """读取模板旁的同名 JSON 档案；无档案或解析失败返回 {}"""
    profile_path = Path(template_path).with_suffix('.json')
    if not profile_path.exists():
        return {}
    try:
        return json.loads(profile_path.read_text(encoding='utf-8-sig'))
    except Exception as exc:
        if warnings is not None:
            warnings.append(f"模板档案 {profile_path.name} 解析失败，按默认规则处理：{exc}")
        return {}


def _resolve_merge_opts(profile: dict) -> tuple[dict, set]:
    """合并档案 merge 段与默认值；返回 (opts, 档案显式声明过的键集合)"""
    declared = set((profile.get('merge') or {}).keys())
    opts = dict(_DEFAULT_MERGE_OPTS)
    opts.update(profile.get('merge') or {})
    # 标签键统一归一化：查找侧（_norm_label）去空白/□/尾冒号，键侧必须同样
    # 归一，否则档案里写「院    系」永远查不到（华中科大封面实测）
    raw_labels = (profile.get('merge') or {}).get('cover_labels') or {}
    labels = dict(_DEFAULT_COVER_LABELS)
    labels.update({_norm_label(k): v for k, v in raw_labels.items()})
    opts['cover_labels'] = labels
    raw_para = (profile.get('merge') or {}).get('paragraph_labels') or {}
    opts['paragraph_labels'] = {
        _norm_label(k): v for k, v in raw_para.items()}
    return opts, declared


# ═══════════════════════════════════════════
# 公共 API
# ═══════════════════════════════════════════

def merge_template(
    out_doc: Document,
    template_path: str | Path,
    metadata: Dict[str, str],
    warnings: Optional[List[str]] = None,
) -> None:
    """将学校模板的封面、说明页、扉页、学术承诺合并到输出文档，
    并把模板末尾的附件（成绩评定表）、校徽封底追加到输出文档结尾。

    切片全部按内容锚点定位（不依赖固定索引，任意结构的学校模板都可用）：
    - 封面起点 = 第一个含「学校代码」的段落（之前的杂项段落忽略）
    - 说明页起点 = 第一个文本为「说明」的段落
    - 前置块终点 = 第一个文本为「摘要」的段落（模板自带示例正文的开头）；
      模板无示例正文时退回第一个段级分节符
    - 尾部 = 最后一个段级分节符之后：「附件N/成绩评定表」段落起是评定表，
      其前是校徽封底画
    锚点缺失时对应部分跳过并记 warning，不中断转换。
    """
    if warnings is None:
        warnings = []
    template_path = Path(template_path)
    if not template_path.exists():
        return

    # 模板格式档案：锚点/标签/占位符规则（无档案则用山财默认）
    profile = load_template_profile(template_path, warnings)
    opts, declared = _resolve_merge_opts(profile)

    tpl_doc = Document(str(template_path))
    tpl_body = tpl_doc.element.body
    children = list(tpl_body)
    n = len(children)

    def norm(i: int) -> str:
        return re.sub(r'\s+', '', _get_element_text(children[i]))

    # 段级分节符位置（body 级 sectPr 不算，它是文档末尾的收束节）
    para_sect_indices = [
        i for i in _find_section_breaks(tpl_body)
        if children[i].tag == qn('w:p')
    ]

    # ── 前置块锚点 ──
    cover_anchor = opts['cover_start_anchor']
    if cover_anchor:
        cover_start = next((i for i in range(n) if cover_anchor in norm(i)), None)
        if cover_start is None:
            cover_start = 0
            if 'cover_start_anchor' not in declared:
                warnings.append(f"模板中未找到「{cover_anchor}」封面锚点，从文档开头取封面")
        elif cover_start > 0 and any(norm(i) for i in range(cover_start)):
            warnings.append(f"模板封面前有 {cover_start} 段多余内容，已忽略")
    else:
        cover_start = 0     # 档案声明：此模板封面就从文档开头开始

    front_anchors = opts['front_end_anchors'] or []
    abstract_idx = next(
        (i for i in range(cover_start + 1, n)
         if any(norm(i) in (a, a + '：') for a in front_anchors)),
        None)
    first_sect = next((i for i in para_sect_indices if i > cover_start), None)
    if abstract_idx is not None:
        front_end = abstract_idx          # 锚点之前全部是前置块（含跨分节的扉页/承诺）
    elif first_sect is not None:
        front_end = first_sect + 1        # 含分节符段落本身（提取时会剥掉 sectPr）
    else:
        front_end = n
        if 'front_end_anchors' not in declared:
            warnings.append("模板中未找到「摘要」或分节符锚点，整个模板按前置块合并")

    expl_anchor = opts['explanation_anchor']
    explanation_idx = None
    if expl_anchor:
        explanation_idx = next(
            (i for i in range(cover_start + 1, front_end) if norm(i) == expl_anchor),
            None)
        if explanation_idx is None and 'explanation_anchor' not in declared:
            warnings.append(f"模板中未找到「{expl_anchor}」页锚点，说明页跳过（可能缺分页）")
    if explanation_idx is None:
        explanation_idx = front_end
    print(f"[template_merge] anchors: cover={cover_start} 说明={explanation_idx} "
          f"front_end={front_end} para_sects={para_sect_indices}")

    # 前置内容全部插入到输出文档第一个元素（封面分节的分节符段落）之前，
    # 使其落在封面分节内——无页眉、无页码，与学校原始文档一致
    out_body = out_doc.element.body
    first_child = _get_first_content_child(out_body)
    rid_map = _build_rid_map(tpl_doc, out_doc)

    cover_elements = _extract_elements(tpl_body, cover_start, explanation_idx, rid_map)
    explanation_elements = _extract_elements(
        tpl_body, explanation_idx, front_end, rid_map)

    # 模板里封面和说明页之间没有显式分页，原模板靠封面内容恰好占满一页自然
    # 分页——论文标题较短时封面变矮，「说 明」会被吸上封面页。
    # 处理：裁掉封面末尾的空段（满页时它们会溢出到下页页首），再给说明首段
    # 加「段前分页」（已在页首时不会产生空白页）。
    while cover_elements and _is_empty_paragraph(cover_elements[-1]):
        cover_elements.pop()
    if explanation_elements:
        _set_page_break_before(explanation_elements[0])

    # 反方向的坑：标题超长时信息表变高，「完成时间」行会被挤到下一页。
    # 封面表格上方有若干装饰性空段，按标题超出的行数裁掉等量空段让位
    # （常规长度标题 extra=0，不影响既有布局；扉页上的同款表格一并处理）
    _title_layout = opts.get('title_layout')
    _labels = opts['cover_labels']
    cover_elements = _compress_spacing_before_tables(
        cover_elements, metadata, _title_layout, _labels, warnings, '封面')
    explanation_elements = _compress_spacing_before_tables(
        explanation_elements, metadata, _title_layout, _labels, warnings, '扉页')

    front_elements = cover_elements + explanation_elements
    for el in cover_elements:
        _fix_signature_underlines(el)
    for el in front_elements:
        if first_child is not None:
            out_body.insert(out_body.index(first_child), el)
        else:
            out_body.append(el)
    print(f"[template_merge] Inserted {len(cover_elements)} cover elements")
    print(f"[template_merge] Inserted {len(explanation_elements)} explanation elements")

    # 封面分节继承模板自己的页面设置，程度由档案 inherit_page_setup 决定：
    #   'frame'（默认）= 纸张 pgSz + 页边距 pgMar
    #   'full'        = 再加行网格 docGrid
    #   'none'        = 全用输出配置
    # 两种需求实测冲突，故必须可声明：
    # - 东南模板封面表格是浮动定位（tblpPr 锚 page），版式同时依赖页框与
    #   行网格，缺 docGrid 时表格会上移盖住标题图 → 该模板声明 'full'。
    # - 附件4 的 linesAndChars/312 网格会把行高吸附到网格，与本项目自己设的
    #   字号/行距叠加后把「承诺+使用授权」挤出一行，连锁多 2 页 → 用默认 'frame'。
    # 页眉页脚引用一律不带（壳分节本就不显示页眉页码）。
    inherit_mode = opts.get('inherit_page_setup') or 'frame'
    inherit_tags = {
        'none': (),
        'frame': ('w:pgSz', 'w:pgMar'),
        'full': ('w:pgSz', 'w:pgMar', 'w:docGrid'),
    }.get(inherit_mode, ('w:pgSz', 'w:pgMar'))
    tpl_first_sect = None
    if para_sect_indices:
        first_sect_p = children[para_sect_indices[0]]
        first_pPr = first_sect_p.find(qn('w:pPr'))
        if first_pPr is not None:
            tpl_first_sect = first_pPr.find(qn('w:sectPr'))
    if tpl_first_sect is None:
        tpl_first_sect = tpl_body.find(qn('w:sectPr'))
    if tpl_first_sect is not None and first_child is not None:
        out_pPr = first_child.find(qn('w:pPr'))
        out_sect = (out_pPr.find(qn('w:sectPr'))
                    if out_pPr is not None else None)
        if out_sect is not None:
            for tag in inherit_tags:
                src = tpl_first_sect.find(qn(tag))
                if src is None:
                    continue
                dst = out_sect.find(qn(tag))
                new = deepcopy(src)
                if dst is not None:
                    out_sect.replace(dst, new)
                else:
                    out_sect.append(new)

    # 追加模板尾部内容，按用户要求的装订顺序重排：
    #   附件2/成绩评定表 → 空白偶数页 → 校徽封底页（全文最后一页，奇数页）
    # 尾部锚点 = 最后一个段级分节符之后：[校徽封底画 ...]['附件N'开头的评定表 ...]
    # 输出文档末尾已由 convert() 预留两个分节：倒数第2节放评定表，最后一节放封底画
    last_sect = para_sect_indices[-1] if para_sect_indices else None
    if (opts['tail'] and last_sect is not None
            and last_sect + 1 < n and last_sect + 1 >= front_end):
        tail_start = last_sect + 1
        form_anchors = opts['tail_form_anchors'] or []
        split_idx = None
        for i in range(tail_start, n):
            t = norm(i)
            if any(t.startswith(a) or (len(a) > 2 and a in t) for a in form_anchors):
                split_idx = i
                break
        if split_idx is None and 'tail_form_anchors' not in declared:
            warnings.append("模板尾部未找到「附件/成绩评定表」锚点，"
                            "尾部内容全部按封底画处理")

        if split_idx is not None:
            back_art = _extract_elements(tpl_body, tail_start, split_idx, rid_map)
            back_form = _extract_elements(tpl_body, split_idx, None, rid_map)
        else:
            back_art = _extract_elements(tpl_body, tail_start, None, rid_map)
            back_form = []

        # 修剪封底画末尾的空段落，避免校徽页后面多出一页空白
        while back_art and _is_empty_paragraph(back_art[-1]):
            back_art.pop()

        # 尾部表单（成绩评定表等）去浮动：浮动表格（tblpPr）按锚点段落偏移定位，
        # 页面几何一变就漂移，再叠加大行高的 cantSplit 行（评语框近 19cm），
        # 整行放不下就被挤到下一页——实测山财毕业论文附件3 指导教师评定表
        # 因此断成「表头一页＋评语框一页」。表单要求整表独立成页且不断裂，
        # 内联随文流才可控。封面/扉页的浮动表格另有版式作用，不在此处理。
        _unfloat_tables(back_form)

        # 锚点：最后一个段落级 sectPr（评定表分节末）和 body 级 sectPr（封底分节）
        last_para_sect = None
        body_sect_pr = None
        for child in out_body:
            if child.tag == qn('w:sectPr'):
                body_sect_pr = child
            elif child.tag == qn('w:p'):
                pPr = child.find(qn('w:pPr'))
                if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
                    last_para_sect = child

        for el in back_form:
            anchor = last_para_sect if last_para_sect is not None else body_sect_pr
            if anchor is not None:
                out_body.insert(out_body.index(anchor), el)
            else:
                out_body.append(el)
        for el in back_art:
            if body_sect_pr is not None:
                out_body.insert(out_body.index(body_sect_pr), el)
            else:
                out_body.append(el)
        print(f"[template_merge] Inserted {len(back_form)} form + {len(back_art)} back-cover elements")

        # 评定表/封底分节同样继承模板尾部原生页框：评定表的巨型评语格
        # 高度按原页边距设计（毕业论文模板上下 2.4cm），套用输出配置的
        # 3.0/2.5cm 会少 0.7cm 可用高度，装不下时整行甩到下一页。
        # 模板尾部内容位于其 body 级 sectPr 所辖分节，取它的页框。
        if inherit_tags:
            tpl_tail_sect = tpl_body.find(qn('w:sectPr'))
            targets = []
            if last_para_sect is not None:
                pPr = last_para_sect.find(qn('w:pPr'))
                if pPr is not None:
                    targets.append(pPr.find(qn('w:sectPr')))
            targets.append(body_sect_pr)
            if tpl_tail_sect is not None:
                for out_sect in targets:
                    if out_sect is None:
                        continue
                    for tag in inherit_tags:
                        src = tpl_tail_sect.find(qn(tag))
                        if src is None:
                            continue
                        dst = out_sect.find(qn(tag))
                        new = deepcopy(src)
                        if dst is not None:
                            out_sect.replace(dst, new)
                        else:
                            out_sect.append(new)
    elif opts['tail']:
        warnings.append("模板中未找到尾部（成绩评定表/封底）内容，已跳过")

    # 填写专业代码
    if opts['major_code_label']:
        major_code = metadata.get("majorCode", "120210")
        _fill_major_code(out_body, major_code, opts['major_code_label'])

    # 填写封面/扉页信息（表格 + 段落式填空 + 正则占位符），并做表格防分页处理
    # （只处理前置内容，不碰尾部成绩评定表；嵌套表也要处理——武大哲学封面
    # 的「姓名 | 张某某」标签/值格在内层表格里，只扫顶层会漏填）
    labels = opts['cover_labels']
    for el in front_elements:
        for tbl_el in el.iter(qn('w:tbl')):
            _fill_cover_table(tbl_el, metadata, labels)
            _fix_cover_table_anti_page_break(tbl_el)
    if opts['paragraph_labels']:
        _fill_cover_paragraphs(front_elements, metadata, opts['paragraph_labels'])
    if opts['placeholders']:
        _apply_placeholders(front_elements, metadata, opts['placeholders'], warnings)


# ═══════════════════════════════════════════
# 内部实现
# ═══════════════════════════════════════════

def _find_section_breaks(body) -> List[int]:
    """查找 body 中所有包含 sectPr 的段落索引"""
    indices = []
    for i, child in enumerate(body):
        # 检查段落属性中的 sectPr
        pPr = child.find(qn('w:pPr'))
        if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
            indices.append(i)
        # 检查 body 级别的 sectPr
        if child.tag == qn('w:sectPr'):
            indices.append(i)
    return indices


def _is_empty_paragraph(el) -> bool:
    """判断元素是否为无文本、无图片的空段落"""
    if el.tag != qn('w:p'):
        return False
    if _get_element_text(el).strip():
        return False
    return next(el.iter(qn('w:drawing')), None) is None


# 标题栏排版参数。正常情况由 _auto_title_layout 从模板自身量出来（栏宽/字号/
# 样例文字），这里只是量不到时的兜底；档案 merge.title_layout 可强制覆盖。
#   title_cols       中文题目栏每行容纳的全角字符当量
#   title_base_lines 模板样例题目占的行数（＝模板原排版的基准）
#   eng_cols / eng_base_lines 同理（英文按半角字符数）
_DEFAULT_TITLE_LAYOUT = {
    'title_cols': 17, 'title_base_lines': 2,
    'eng_cols': 34, 'eng_base_lines': 4,
}

_TWIPS_PER_PT = 20
_DEFAULT_CELL_MARGIN_TW = 108      # Word 默认单元格左右边距 0.19cm
_LATIN_WIDTH_RATIO = 0.5           # Times New Roman 平均字宽 ≈ 0.5 em


def _cell_text_width_tw(tbl, tc) -> Optional[int]:
    """单元格的可排字宽度（twips）＝ 栏宽 − 左右内边距；宽度非绝对值时返回 None"""
    tcPr = tc.find(qn('w:tcPr'))
    tcW = tcPr.find(qn('w:tcW')) if tcPr is not None else None
    if tcW is None:
        return None
    if (tcW.get(qn('w:type')) or 'dxa') != 'dxa':
        return None
    try:
        width = int(tcW.get(qn('w:w')))
    except (TypeError, ValueError):
        return None
    tblPr = tbl.find(qn('w:tblPr'))
    mar = tblPr.find(qn('w:tblCellMar')) if tblPr is not None else None
    pad = 0
    for side in ('w:left', 'w:right'):
        node = mar.find(qn(side)) if mar is not None else None
        try:
            pad += int(node.get(qn('w:w'))) if node is not None else _DEFAULT_CELL_MARGIN_TW
        except (TypeError, ValueError):
            pad += _DEFAULT_CELL_MARGIN_TW
    return max(1, width - pad)


def _cell_font_size_tw(tc, default_half_pt: int = 28) -> float:
    """单元格首个 run 的字号换算成 twips（＝全角字宽）"""
    for sz in tc.iter(qn('w:sz')):
        try:
            return int(sz.get(qn('w:val'))) / 2 * _TWIPS_PER_PT
        except (TypeError, ValueError):
            break
    return default_half_pt / 2 * _TWIPS_PER_PT


def _disp_width(text: str) -> float:
    """显示宽度：全角 1、半角 0.5（「AIGC」这类混排按字符数会高估行数）"""
    return sum(1.0 if ord(c) > 0x2E80 else 0.5 for c in text)


def _auto_title_layout(tbl, labels: Dict[str, str]) -> dict:
    """从模板表格自身量出题目栏每行字数与样例占的行数。

    这样换任何模板都不需要人工标定：栏宽（tcW − 单元格内边距）除以字号
    即每行全角字数，模板里原有的样例题目文字则给出基准行数。
    """
    import math
    out: dict = {}
    for row in tbl.findall(qn('w:tr')):
        cells = row.findall(qn('w:tc'))
        for i, cell in enumerate(cells[:-1]):
            key = labels.get(_norm_label(_get_element_text(cell)))
            if key not in ('title', 'engTitle'):
                continue
            # 值格定位规则与 _fill_cover_table 保持一致（跳过纯冒号格）
            target = None
            for j in range(i + 1, len(cells)):
                text = _get_element_text(cells[j]).strip()
                if _norm_label(text) in labels:
                    break
                if text in ('：', ':'):
                    continue
                target = cells[j]
                break
            if target is None:
                continue
            width = _cell_text_width_tw(tbl, target)
            if not width:
                continue
            char_tw = _cell_font_size_tw(target)
            if char_tw <= 0:
                continue
            sample = _get_element_text(target).strip()
            if key == 'title':
                cols = max(1, int(width / char_tw))
                out['title_cols'] = cols
                out['title_base_lines'] = (
                    max(1, math.ceil(_disp_width(re.sub(r'\s+', '', sample)) / cols))
                    if sample else 1)
            else:
                cols = max(1, int(width / (char_tw * _LATIN_WIDTH_RATIO)))
                out['eng_cols'] = cols
                out['eng_base_lines'] = (
                    max(1, math.ceil(len(sample) / cols)) if sample else 1)
    return out


def _table_extra_lines(tbl, metadata: Dict[str, str],
                       layout: Optional[dict],
                       labels: Dict[str, str]) -> int:
    """估算某张信息表因题目变长而超出模板样例基准的行数。

    **必须按表分别算**：封面表常只有中文题目，扉页表中英文题目都有，
    用同一个总数会让封面裁多、扉页裁少。
    中文按显示宽度折算（全角=1、半角=0.5，否则「AIGC」这类混排会高估行数），
    英文按字符数。
    """
    import math
    # 优先用从模板实测的参数，档案里显式给的再覆盖它（人工兜底）
    lo = dict(_DEFAULT_TITLE_LAYOUT)
    lo.update(_auto_title_layout(tbl, labels))
    lo.update(layout or {})

    present = set()
    for row in tbl.findall(qn('w:tr')):
        for cell in row.findall(qn('w:tc')):
            key = labels.get(_norm_label(_get_element_text(cell)))
            if key in ('title', 'engTitle'):
                present.add(key)

    extra = 0
    title = metadata.get('title', '')
    if 'title' in present and title:
        extra += max(0, math.ceil(_disp_width(title) / lo['title_cols'])
                     - lo['title_base_lines'])
    eng = metadata.get('engTitle', '')
    if 'engTitle' in present and eng:
        extra += max(0, math.ceil(len(eng) / lo['eng_cols']) - lo['eng_base_lines'])
    return extra


def _table_compact_savings(tbl, metadata: Dict[str, str],
                           layout: Optional[dict],
                           labels: Dict[str, str]) -> int:
    """估算题目栏由 1.5 倍行距收紧为单倍后省出的行数（每行省半行）"""
    import math
    lo = dict(_DEFAULT_TITLE_LAYOUT)
    lo.update(_auto_title_layout(tbl, labels))
    lo.update(layout or {})

    present = set()
    for row in tbl.findall(qn('w:tr')):
        for cell in row.findall(qn('w:tc')):
            key = labels.get(_norm_label(_get_element_text(cell)))
            if key in ('title', 'engTitle'):
                present.add(key)

    saved = 0.0
    title = metadata.get('title', '')
    if 'title' in present and title:
        lines = math.ceil(_disp_width(title) / lo['title_cols'])
        if lines > lo['title_base_lines']:
            saved += lines * 0.5
    eng = metadata.get('engTitle', '')
    if 'engTitle' in present and eng:
        lines = math.ceil(len(eng) / lo['eng_cols'])
        if lines > lo['eng_base_lines']:
            saved += lines * 0.5
    return int(saved)


def _compress_spacing_before_tables(elements: List, metadata: Dict[str, str],
                                    layout: Optional[dict],
                                    labels: Dict[str, str],
                                    warnings: Optional[List[str]] = None,
                                    where: str = '封面') -> List:
    """按各表格自身的超出行数，删除其紧邻前置空段以腾出竖向空间。

    空段不够抵消时（超长题目）记 warning——此时题目栏已自动收紧行距，
    仍可能把末行挤到下一页，需要人工缩短题目或调模板。
    """
    drop = set()
    for idx, el in enumerate(elements):
        if el.tag != qn('w:tbl'):
            continue
        count = _table_extra_lines(el, metadata, layout, labels)
        if count <= 0:
            continue
        removed = 0
        j = idx - 1
        while j >= 0 and removed < count and _is_empty_paragraph(elements[j]):
            drop.add(j)
            removed += 1
            j -= 1
        # 收紧行距（1.5→1.0 倍）本身也能省出高度，先折抵再判断是否真的不够
        saved = _table_compact_savings(el, metadata, layout, labels)
        if removed + saved < count and warnings is not None:
            msg = (f"题目过长：{where}信息表比模板样例高出约 {count} 行，"
                   f"裁空段腾出 {removed} 行、收紧题目栏行距折抵 {saved} 行，"
                   f"仍差约 {count - removed - saved} 行；末行可能被挤到下一页，"
                   f"建议缩短题目或改用题目栏更宽的模板")
            if msg not in warnings:
                warnings.append(msg)
    return [el for i, el in enumerate(elements) if i not in drop]


def _unfloat_tables(elements: List) -> None:
    """把元素里的浮动表格（tblpPr）改成内联表格（就地修改）"""
    for el in elements:
        for tbl in el.iter(qn('w:tbl')):
            tblPr = tbl.find(qn('w:tblPr'))
            if tblPr is None:
                continue
            tblpPr = tblPr.find(qn('w:tblpPr'))
            if tblpPr is not None:
                tblPr.remove(tblpPr)


def _set_page_break_before(el) -> None:
    """给段落加 w:pageBreakBefore（段前分页；段落已在页首时不产生空白页）"""
    if el.tag != qn('w:p'):
        return
    pPr = el.find(qn('w:pPr'))
    if pPr is None:
        pPr = el.makeelement(qn('w:pPr'), {})
        el.insert(0, pPr)
    if pPr.find(qn('w:pageBreakBefore')) is not None:
        return
    pbb = pPr.makeelement(qn('w:pageBreakBefore'), {})
    # schema 顺序：pageBreakBefore 需在 pStyle/keepNext/keepLines 之后
    anchor = None
    for tag in ('w:pStyle', 'w:keepNext', 'w:keepLines'):
        found = pPr.find(qn(tag))
        if found is not None:
            anchor = found
    if anchor is not None:
        anchor.addnext(pbb)
    else:
        pPr.insert(0, pbb)


def _get_first_content_child(body):
    """获取 body 中第一个非 sectPr 的子元素"""
    for child in body:
        if child.tag != qn('w:sectPr'):
            return child
    return None


# 复制模板内容时必须剥掉的跨部件引用：目标部件（批注/脚注/尾注/OLE 嵌入件）
# 不会被合并进输出文档，悬空引用会让 Word 报「文件可能已经损坏」拒绝打开
# （实测：东南大学模板的 AI 说明表带一条批注；川大完整版封面校徽是
# StaticMetafile OLE 对象——剥掉 o:OLEObject 后 v:imagedata 静态图照常显示）
_FOREIGN_REF_TAGS = (
    'w:commentRangeStart', 'w:commentRangeEnd', 'w:commentReference',
    'w:footnoteReference', 'w:endnoteReference',
)
_OLE_OBJECT_TAG = '{urn:schemas-microsoft-com:office:office}OLEObject'


def _strip_foreign_refs(el) -> None:
    """删除元素内的批注/脚注/尾注引用与 OLE 对象节点（引用目标不随切片复制）"""
    for tag in _FOREIGN_REF_TAGS:
        for node in list(el.iter(qn(tag))):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
    for node in list(el.iter(_OLE_OBJECT_TAG)):
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)
    # 剥掉 OLE 后的 w:object 转成普通 w:pict：o:ole/spid 标记留着时 Word
    # 视其为缺载荷的 OLE，报「文件可能已经损坏」（华中科大 AI 学院封面横幅实测）
    O_NS = '{urn:schemas-microsoft-com:office:office}'
    for obj in list(el.iter(qn('w:object'))):
        pict = obj.makeelement(qn('w:pict'), {})
        for child in list(obj):
            obj.remove(child)
            pict.append(child)
        for shape in pict.iter('{urn:schemas-microsoft-com:vml}shape'):
            shape.attrib.pop(O_NS + 'ole', None)
            shape.attrib.pop(O_NS + 'spid', None)
        parent = obj.getparent()
        if parent is not None:
            parent.replace(obj, pict)


def _extract_elements(body, start_idx: int, end_idx: Optional[int], rid_map: Optional[Dict[str, str]] = None) -> List:
    """提取 body 中指定范围的元素（深度拷贝，去除 sectPr，保留并修正图片引用）"""
    children = list(body)
    if end_idx is None:
        end_idx = len(children)
    end_idx = min(end_idx, len(children))

    elements = []
    for i in range(start_idx, end_idx):
        child = children[i]
        if child.tag == qn('w:sectPr'):
            continue
        el = deepcopy(child)
        _strip_foreign_refs(el)
        # 移除段落中的 sectPr
        pPr = el.find(qn('w:pPr'))
        if pPr is not None:
            sect_pr = pPr.find(qn('w:sectPr'))
            if sect_pr is not None:
                pPr.remove(sect_pr)
                # 范围末尾的分节符不补分页符（后续内容自带分节符换页，避免多出空白页）
                if i < end_idx - 1:
                    run = el.makeelement(qn('w:r'), {})
                    br = run.makeelement(qn('w:br'), {qn('w:type'): 'page'})
                    run.append(br)
                    el.append(run)
        # 更新图片引用的 rId，使其指向输出文档中的图片
        if rid_map:
            _update_image_refs(el, rid_map)
        elements.append(el)
    return elements


def _build_rid_map(src_doc: Document, dst_doc: Document) -> Dict[str, str]:
    """将模板文档中的图片复制到输出文档，返回 {old_rId: new_rId} 映射。

    必须按字节拷贝并由输出文档分配新部件名（get_or_add_image）：
    直接 relate_to 模板包里的图片部件会与输出文档自己 add_picture 产生的
    /word/media/imageN 部件名冲突，保存出的 zip 含重名条目，Word 拒绝打开。

    get_or_add_image 只认 png/jpg 等常规格式；OLE 预览图（WMF/EMF 静态
    元文件，川大完整版/华中科大 AI 学院封面校徽）会抛异常——这类图改为
    手工建部件按字节拷贝，否则映射缺失、v:imagedata 悬空指向输出文档的
    其它部件（实测指向 theme），Word 报「文件可能已经损坏」。
    """
    rid_map: Dict[str, str] = {}
    dst_part = dst_doc.part
    for rel_id, rel in src_doc.part.rels.items():
        if rel.reltype != RT.IMAGE:
            continue
        try:
            new_rid, _ = dst_part.get_or_add_image(io.BytesIO(rel.target_part.blob))
            rid_map[rel_id] = new_rid
            continue
        except Exception:
            pass
        try:
            blob = rel.target_part.blob
            content_type = rel.target_part.content_type
            suffix = Path(str(rel.target_part.partname)).suffix or ".bin"
            n = 1
            while True:
                partname = PackURI(f"/word/media/template_image{n}{suffix}")
                if not any(p.partname == partname
                           for p in dst_part.package.iter_parts()):
                    break
                n += 1
            part = Part(partname, content_type, blob, dst_part.package)
            new_rid = dst_part.relate_to(part, RT.IMAGE)
            rid_map[rel_id] = new_rid
        except Exception:
            pass
    return rid_map


def _update_image_refs(element, rid_map: Dict[str, str]) -> None:
    """更新 XML 元素中所有图片引用的 rId。

    两种图片引用都要处理：
    - DrawingML：a:blip r:embed（现代 Word 插图）
    - VML：v:imagedata r:id（旧式图片，中财封面校徽实测为此类）——
      不修正会指向输出文档里不相干的部件，Word 报「文件可能已经损坏」
    """
    nsmap_r = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    nsmap_o = '{urn:schemas-microsoft-com:office:office}'
    for blip in element.iter(qn('a:blip')):
        old_rid = blip.get(nsmap_r + 'embed')
        if old_rid and old_rid in rid_map:
            blip.set(nsmap_r + 'embed', rid_map[old_rid])
    for node in element.iter():
        if isinstance(node.tag, str) and node.tag.endswith('}imagedata'):
            old_rid = node.get(nsmap_r + 'id')
            if old_rid and old_rid in rid_map:
                node.set(nsmap_r + 'id', rid_map[old_rid])
            old_relid = node.get(nsmap_o + 'relid')
            if old_relid and old_relid in rid_map:
                node.set(nsmap_o + 'relid', rid_map[old_relid])


def _fix_signature_underlines(element) -> None:
    """修复签名处的下划线长度（对应 C# FixSignatureUnderlines）"""
    for p in element.iter(qn('w:p')):
        p_text = _get_element_text(p)
        if not any(kw in p_text for kw in ('签名', '期', '指导教师')):
            continue

        for r in p.findall(qn('w:r')):
            rPr = r.find(qn('w:rPr'))
            if rPr is None:
                continue
            u = rPr.find(qn('w:u'))
            if u is None:
                continue
            t = r.find(qn('w:t'))
            if t is None or t.text is None:
                continue

            text = t.text
            if text.strip() == '' and len(text) > 0:
                t.text = ' ' * max(24, len(text) * 2)
                t.set(qn('xml:space'), 'preserve')
            elif (text.strip() and ' ' in text
                  and len(text) - len(text.rstrip()) > len(text.strip())):
                trimmed = text.rstrip()
                space_count = len(text) - len(trimmed)
                t.text = trimmed + ' ' * max(24, space_count * 2)
                t.set(qn('xml:space'), 'preserve')


def _fill_major_code(body, major_code: str, label: str = '专业代码') -> None:
    """填写专业代码（对应 C# FillMajorCode）"""
    for p in body.iter(qn('w:p')):
        p_text = _get_element_text(p)
        if label not in p_text:
            continue

        runs = p.findall(qn('w:r'))
        for i, r in enumerate(runs):
            t = r.find(qn('w:t'))
            if t is None or t.text is None or label not in t.text:
                continue

            # 标签之后第一个带下划线的 run 才是填空线（Word 规范化可能在
            # 标签和填空线之间插入空格 run，不能简单取「下一个」）
            target_t = None
            for nxt in runs[i + 1:]:
                rPr = nxt.find(qn('w:rPr'))
                nt = nxt.find(qn('w:t'))
                if nt is None:
                    continue
                if rPr is not None and rPr.find(qn('w:u')) is not None:
                    target_t = nt
                    break
                if target_t is None:
                    target_t = nt   # 兜底：记住第一个有文本的 run
            if target_t is None:
                break

            # 居中填入专业代码（总宽度15字符）
            total_width = 15
            pad_left = (total_width - len(major_code)) // 2
            pad_right = total_width - len(major_code) - pad_left
            target_t.text = ' ' * pad_left + major_code + ' ' * pad_right
            target_t.set(qn('xml:space'), 'preserve')
            break


_META_DEFAULTS = {
    'title': '论文标题',
    'engTitle': 'Paper Title',
    'name': '姓名',
    'studentId': '学号',
    'class': '班级',
    'major': '专业',
    'college': '学院',
    'advisor': '指导教师',
    'date': '2026年5月',
}


def _norm_label(text: str) -> str:
    """标签归一化：去所有空白、去 □ 间隔符、去尾部冒号（「学    号：」→「学号」）。

    □（U+25A1）是哈工大等模板里用来拉开字距的字面字符（「摘□□要」「答□辩□日□期」），
    与空白一样只承担排版作用，不参与标签语义。
    """
    return re.sub(r'[\s\u25a1]+', '', text).rstrip('：:')


def _fill_cover_table(tbl, metadata: Dict[str, str],
                      labels: Dict[str, str]) -> None:
    """填写封面/扉页信息表字段（对应 C# FillCoverTable）。

    通用单元格规则（兼容三种实测排布）：
    - 山财  [标签 | ： | 值]（3 格，中间纯冒号格跳过）
    - 东南  [标签 | 值]（2 格）
    - AI 表 [标签 | 值 | 标签 | 值]（4 格，一行多组）
    逐格扫描：标签格右侧第一个「既非标签、亦非纯冒号」的格子就是值格。
    题目超出模板样例行数时，该栏自动由 1.5 倍行距收紧为单倍行距，
    减少把表格末行（完成时间等）挤到下一页的概率。
    """
    import math
    lo = dict(_DEFAULT_TITLE_LAYOUT)
    lo.update(_auto_title_layout(tbl, labels))

    def _needs_compact(meta_key: str, value: str) -> bool:
        if not value:
            return False
        if meta_key == 'title':
            lines = math.ceil(_disp_width(value) / lo['title_cols'])
            return lines > lo['title_base_lines']
        if meta_key == 'engTitle':
            lines = math.ceil(len(value) / lo['eng_cols'])
            return lines > lo['eng_base_lines']
        return False

    for row in tbl.findall(qn('w:tr')):
        cells = row.findall(qn('w:tc'))
        i = 0
        while i < len(cells):
            label_norm = _norm_label(_get_element_text(cells[i]))
            meta_key = labels.get(label_norm)
            if meta_key is None:
                i += 1
                continue
            # 右侧找值格：跳过纯冒号/空白连接格中的冒号格
            target_cell = None
            j = i + 1
            while j < len(cells):
                cand_norm = _norm_label(_get_element_text(cells[j]))
                if cand_norm in labels:
                    break                      # 撞到下一个标签，本组无值格
                if _get_element_text(cells[j]).strip() in ('：', ':'):
                    j += 1
                    continue
                target_cell = cells[j]
                break
            if target_cell is None:
                i += 1
                continue

            is_english = (meta_key == 'engTitle')
            value = metadata.get(meta_key, '') or _META_DEFAULTS.get(meta_key, '')
            _write_cover_cell(target_cell, value, is_english,
                              compact=_needs_compact(meta_key, value))
            i = j + 1
        # 行内继续扫描由 while 完成


def _write_cover_cell(target_cell, value: str, is_english: bool,
                      compact: bool = False) -> None:
        """清空单元格并按封面样式写入值（compact＝题目超长，收紧行距）"""
        # 移除现有段落
        for old_p in target_cell.findall(qn('w:p')):
            target_cell.remove(old_p)

        # 创建新段落
        p = target_cell.makeelement(qn('w:p'), {})
        pPr = p.makeelement(qn('w:pPr'), {})
        spacing = pPr.makeelement(qn('w:spacing'), {
            qn('w:line'): '240' if compact else '360',
            qn('w:lineRule'): 'auto',
        })
        pPr.append(spacing)
        jc = pPr.makeelement(qn('w:jc'), {
            qn('w:val'): 'left' if is_english else 'center',
        })
        pPr.append(jc)
        p.append(pPr)

        if is_english:
            # 英文题目：Times New Roman 斜体
            run = p.makeelement(qn('w:r'), {})
            rPr = run.makeelement(qn('w:rPr'), {})
            rFonts = rPr.makeelement(qn('w:rFonts'), {
                qn('w:ascii'): 'Times New Roman',
                qn('w:hAnsi'): 'Times New Roman',
                qn('w:eastAsia'): 'Times New Roman',
            })
            rPr.append(rFonts)
            italic = rPr.makeelement(qn('w:i'), {})
            rPr.append(italic)
            sz = rPr.makeelement(qn('w:sz'), {qn('w:val'): '28'})
            rPr.append(sz)
            run.append(rPr)
            t = run.makeelement(qn('w:t'), {})
            t.text = value
            t.set(qn('xml:space'), 'preserve')
            run.append(t)
            p.append(run)
        else:
            # 中文字段：区分中英文字符，分别设置字体
            _add_mixed_run(p, value)

        target_cell.append(p)


def _fill_cover_paragraphs(elements: List, metadata: Dict[str, str],
                           labels: Dict[str, str]) -> None:
    """段落式封面填空：「学生姓名：__x x x__」这类标签+下划线占位段。

    识别：段落归一化文本以「标签：」开头 → 冒号之后的 run 全部视为填空区，
    值居中写入第一个有文本的 run（保留其下划线等格式），其余 run 清空。
    """
    for el in elements:
        for p in el.iter(qn('w:p')):
            full = re.sub(r'\s+', '', _get_element_text(p))
            meta_key = None
            for label, key in labels.items():
                if full.startswith(label + '：') or full.startswith(label + ':'):
                    meta_key = key
                    break

            # 找到标签冒号所在的 run（标签可能被拆成多个 run）
            runs = p.findall(qn('w:r'))
            acc = ''
            colon_run_idx = None
            for idx, r in enumerate(runs):
                t = r.find(qn('w:t'))
                if t is not None and t.text:
                    acc += re.sub(r'\s+', '', t.text)
                if '：' in acc or ':' in acc:
                    colon_run_idx = idx
                    break

            if meta_key is None:
                # 无冒号变体：「学    院 ____」「题 目_____」标签段（川大完整版/
                # 西北大学/南大封面）——非下划线 run 拼出的文本以标签开头时，
                # 值依次写入填空槽（下划线格或纯 '_' 字符 run，保留填空线）。
                # 连续的下划线 run 归并为一个槽（值写首 run、清余 run）——
                # 矿大扉页「作 者 __某某某__ 学 号 __000000__」一格双栏，槽按
                # 下划线连续段划分而不是按 run 数。
                # 标签可多段连续匹配（南大「年级学号」、矿大「作者学号」），
                # 每段消耗一个槽。
                label_text = ''
                fill_groups = []           # [ [t, t, ...], ... ] 连续填空 run 组
                last_slot_idx = None
                for r_i, r in enumerate(runs):
                    rpr = r.find(qn('w:rPr'))
                    has_u = rpr is not None and rpr.find(qn('w:u')) is not None
                    t = r.find(qn('w:t'))
                    txt = t.text if (t is not None and t.text) else ''
                    is_slot = bool(txt) and (
                        re.fullmatch(r'_+', txt.strip()) or has_u)
                    if is_slot:
                        if last_slot_idx is not None and r_i == last_slot_idx + 1 \
                                and fill_groups:
                            fill_groups[-1].append(t)   # 连续下划线 → 同槽
                        else:
                            fill_groups.append([t])
                        last_slot_idx = r_i
                    elif not has_u:
                        label_text += txt
                        if txt:
                            last_slot_idx = None
                label_norm = re.sub(r'\s+', '', label_text)
                slot_i = 0
                while label_norm and slot_i < len(fill_groups):
                    hit = None
                    for label, key in sorted(labels.items(),
                                             key=lambda kv: -len(kv[0])):
                        if label_norm.startswith(label):
                            hit = (label, key)
                            break
                    if hit is None:
                        break
                    label, key = hit
                    label_norm = label_norm[len(label):]
                    value = (metadata.get(key, '')
                             or _META_DEFAULTS.get(key, ''))
                    group = fill_groups[slot_i]
                    group[0].text = value
                    group[0].set(qn('xml:space'), 'preserve')
                    for t in group[1:]:
                        t.text = ''
                    slot_i += 1
                # 剩余填空组清空（未匹配到标签的空段不留样张值）
                if slot_i:
                    for group in fill_groups[slot_i:]:
                        for t in group:
                            t.text = ''
                elif fill_groups:
                    continue
                elif re.search(r'_+$', label_norm):
                    # 标签与字面下划线同 run（「学生姓名____」，西工大）：
                    # 把段内末尾的下划线串替换为值
                    for label, key in sorted(labels.items(),
                                             key=lambda kv: -len(kv[0])):
                        if not label_norm.startswith(label):
                            continue
                        value = (metadata.get(key, '')
                                 or _META_DEFAULTS.get(key, ''))
                        for r in reversed(runs):
                            t = r.find(qn('w:t'))
                            if t is None or not t.text or not re.search(r'_+$', t.text):
                                continue
                            t.text = re.sub(r'_+$', value, t.text, count=1)
                            t.set(qn('xml:space'), 'preserve')
                            break
                        break
                continue

            if colon_run_idx is None:
                # 冒号前缀命中但 run 里找不到冒号（冒号在文本框等处）→ 不动
                continue
            value = metadata.get(meta_key, '') or _META_DEFAULTS.get(meta_key, '')

            # 冒号之后：值（前后各补 2 个空格）写入第一个 run，其余 run 删除。
            # 不按原下划线宽度撑满——中文值实际字宽大于空格估计，原行又常常
            # 已接近排满，撑满必换行（中财封面指导教师/日期行实测）
            fill_runs = [r for r in runs[colon_run_idx + 1:]
                         if r.find(qn('w:t')) is not None]
            if not fill_runs:
                # 冒号后没有独立的文本 run：要么冒号与样张值同 run
                # （「导师：教授」，值替换冒号后文字），要么是纯空行
                # （「作 者：」，矿大封面）——新建 run 承接值并沿用末 run 格式
                colon_t = runs[colon_run_idx].find(qn('w:t'))
                if colon_t is not None and colon_t.text:
                    m = re.search(r'[：:]', colon_t.text)
                    after = (colon_t.text[m.end():] or '').strip()
                    if after:
                        colon_t.text = colon_t.text[:m.end()] + value
                        colon_t.set(qn('xml:space'), 'preserve')
                        continue
                # 标签行独立成段、值在下一段（武大哲学封面「姓 名：」/「张某某」）：
                # 值写入下一段，不在标签行追加
                if _fill_value_in_next_paragraph(p, labels, value):
                    continue
                new_r = p.makeelement(qn('w:r'), {})
                if runs:
                    last_rpr = runs[-1].find(qn('w:rPr'))
                    if last_rpr is not None:
                        new_r.append(deepcopy(last_rpr))
                new_t = new_r.makeelement(qn('w:t'), {})
                new_t.text = '  ' + value + '  '
                new_t.set(qn('xml:space'), 'preserve')
                new_r.append(new_t)
                p.append(new_r)
                continue
            first_t = fill_runs[0].find(qn('w:t'))
            first_t.text = '  ' + value + '  '
            first_t.set(qn('xml:space'), 'preserve')
            for r in fill_runs[1:]:
                p.remove(r)


def _fill_value_in_next_paragraph(p, labels: Dict[str, str], value: str) -> bool:
    """「标签：」行 + 下一行独立值段（武大哲学封面）时，把值写进下一段。

    条件：标签行除标签外无别的文字；下一段是同父级的短文本段、自身不是
    标签行。命中返回 True 并完成写入。
    """
    if value == '':
        return False
    parent = p.getparent()
    if parent is None:
        return False
    siblings = list(parent)
    try:
        idx = siblings.index(p)
    except ValueError:
        return False
    if idx + 1 >= len(siblings):
        return False
    nxt = siblings[idx + 1]
    if nxt.tag != qn('w:p'):
        return False
    nxt_text = _get_element_text(nxt).strip()
    norm_text = re.sub(r'\s+', '', nxt_text)
    if not norm_text or len(norm_text) > 20:
        return False
    if '：' in norm_text or ':' in norm_text:
        return False
    for label in labels:
        if norm_text.startswith(label):
            return False
    nxt_runs = [r for r in nxt.findall(qn('w:r')) if r.find(qn('w:t')) is not None]
    if not nxt_runs:
        return False
    first_t = nxt_runs[0].find(qn('w:t'))
    first_t.text = value
    first_t.set(qn('xml:space'), 'preserve')
    for r in nxt_runs[1:]:
        t = r.find(qn('w:t'))
        if t is not None:
            t.text = ''
    return True


def _apply_placeholders(elements: List, metadata: Dict[str, str],
                        placeholders: List[dict],
                        warnings: List[str]) -> None:
    """正则占位符替换：整段文本匹配 pattern 时，用元数据值（或空串）替换。

    典型用法：中财封面标题行「XXXX…X」→ {"pattern": "^X{6,}$", "value": "title"}；
    东南标题表提示行「（如此行空白，可自行删除）」→ {"pattern": "…", "value": ""}。
    值写入第一个有文本的 run（继承其字体格式），其余 run 清空。
    """
    compiled = []
    for ph in placeholders:
        try:
            compiled.append((re.compile(ph.get('pattern', '')), ph.get('value', '')))
        except re.error as exc:
            warnings.append(f"模板档案占位符正则无效（{ph.get('pattern')}）：{exc}")
    for el in elements:
        for p in el.iter(qn('w:p')):
            full = re.sub(r'\s+', '', _get_element_text(p))
            if not full:
                continue
            for pattern, value_key in compiled:
                if not pattern.match(full):
                    continue
                value = ('' if value_key == ''
                         else metadata.get(value_key, '')
                         or _META_DEFAULTS.get(value_key, ''))
                wrote = False
                for r in p.findall(qn('w:r')):
                    t = r.find(qn('w:t'))
                    if t is None:
                        continue
                    if not wrote and value:
                        t.text = value
                        t.set(qn('xml:space'), 'preserve')
                        wrote = True
                    else:
                        t.text = ''
                break


def _add_mixed_run(parent, text: str) -> None:
    """为混合中英文的文本创建 run，中文用宋体，英文/数字用 Times New Roman"""
    # 按字符类型分组
    segments = []
    current = ''
    is_latin = None

    for ch in text:
        ch_is_latin = (
            ('0' <= ch <= '9') or ('a' <= ch <= 'z') or ('A' <= ch <= 'Z')
            or ch in '/-. '
        )
        if ch == ' ':
            ch_is_latin = is_latin  # 空格跟随前一个字符类型

        if is_latin is None:
            is_latin = ch_is_latin
            current = ch
        elif ch_is_latin == is_latin:
            current += ch
        else:
            segments.append((current, is_latin))
            current = ch
            is_latin = ch_is_latin

    if current:
        segments.append((current, is_latin))

    for seg_text, seg_is_latin in segments:
        run = parent.makeelement(qn('w:r'), {})
        rPr = run.makeelement(qn('w:rPr'), {})
        bold = rPr.makeelement(qn('w:b'), {})
        rPr.append(bold)
        sz = rPr.makeelement(qn('w:sz'), {qn('w:val'): '28'})
        rPr.append(sz)
        szCs = rPr.makeelement(qn('w:szCs'), {qn('w:val'): '28'})
        rPr.append(szCs)

        if seg_is_latin:
            rFonts = rPr.makeelement(qn('w:rFonts'), {
                qn('w:ascii'): 'Times New Roman',
                qn('w:hAnsi'): 'Times New Roman',
                qn('w:eastAsia'): 'Times New Roman',
            })
        else:
            rFonts = rPr.makeelement(qn('w:rFonts'), {
                qn('w:ascii'): 'Times New Roman',
                qn('w:hAnsi'): 'Times New Roman',
                qn('w:eastAsia'): 'SimSun',
            })
        rPr.append(rFonts)
        run.append(rPr)

        t = run.makeelement(qn('w:t'), {})
        t.text = seg_text
        t.set(qn('xml:space'), 'preserve')
        run.append(t)
        parent.append(run)


def _get_element_text(element) -> str:
    """递归获取元素的全部文本内容"""
    texts = []
    for t in element.iter(qn('w:t')):
        if t.text:
            texts.append(t.text)
    return ''.join(texts)


def _fix_cover_table_anti_page_break(tbl) -> None:
    """封面表格防分页处理：禁止行跨页断行、删除段落分页约束"""
    # 遍历所有行
    for row in tbl.findall(qn('w:tr')):
        # 获取或创建行属性
        trPr = row.find(qn('w:trPr'))
        if trPr is None:
            trPr = row.makeelement(qn('w:trPr'), {})
            row.insert(0, trPr)
        
        # 禁止跨页断行（cantSplit）
        if trPr.find(qn('w:cantSplit')) is None:
            cantSplit = trPr.makeelement(qn('w:cantSplit'), {})
            trPr.append(cantSplit)
        
        # 行高设为最小值（删除固定行高）
        trHeight = trPr.find(qn('w:trHeight'))
        if trHeight is not None:
            hRule = trHeight.get(qn('w:hRule'))
            if hRule == 'exact':
                # 改为 atLeast
                trHeight.set(qn('w:hRule'), 'atLeast')
        
        # 遍历单元格内的段落，删除分页约束
        for cell in row.findall(qn('w:tc')):
            for para in cell.findall(qn('w:p')):
                pPr = para.find(qn('w:pPr'))
                if pPr is None:
                    continue
                # 删除段前分页、段中不分页、与下段同页
                for tag in ['w:pageBreakBefore', 'w:keepLines', 'w:keepNext']:
                    elem = pPr.find(qn(tag))
                    if elem is not None:
                        pPr.remove(elem)
    
    print("[template_merge] Applied anti-page-break settings to cover table")
