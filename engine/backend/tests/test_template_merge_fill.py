# -*- coding: utf-8 -*-
"""_fill_cover_paragraphs 无冒号下划线填空回退（川大完整版封面形态）。"""
from docx import Document
from docx.oxml.ns import qn

from backend.thesis_builder.template_merge import _fill_cover_paragraphs


def _cover_paragraph(runs_spec):
    """runs_spec: [(text, underlined), ...] → 返回段落 XML 元素"""
    doc = Document()
    p = doc.add_paragraph()
    for text, underlined in runs_spec:
        run = p.add_run(text)
        if underlined:
            run.font.underline = True
    return p._p


def test_no_colon_label_fills_first_underlined_run():
    # 「学    院 ____计算机学院____」：标签无冒号，值区为下划线 run
    p = _cover_paragraph([
        ('学', False), ('    ', False), ('院', False), (' ', False),
        ('          ', True), ('计算机学院', True), ('            ', True),
    ])
    _fill_cover_paragraphs([p], {'college': '测试学院'}, {'学院': 'college'})
    texts = [(r.find(qn('w:t')).text if r.find(qn('w:t')) is not None else None)
             for r in p.findall(qn('w:r'))]
    assert texts[4] == '测试学院'
    assert not texts[5] and not texts[6]


def test_colon_form_still_uses_colon_branch():
    # 有冒号时走原冒号分支：值写入冒号后第一个 run，其余 run 删除
    p = _cover_paragraph([('学生姓名', False), ('：', False), ('张三', False)])
    _fill_cover_paragraphs([p], {'name': '李四'}, {'学生姓名': 'name'})
    texts = [(r.find(qn('w:t')).text if r.find(qn('w:t')) is not None else None)
             for r in p.findall(qn('w:r'))]
    assert texts[2] == '  李四  '


def test_no_match_leaves_paragraph_untouched():
    p = _cover_paragraph([('教务处制表', False)])
    original = ''.join(r.find(qn('w:t')).text for r in p.findall(qn('w:r')))
    _fill_cover_paragraphs([p], {'name': '李四'}, {'学生姓名': 'name'})
    assert ''.join(r.find(qn('w:t')).text for r in p.findall(qn('w:r'))) == original


def test_no_colon_without_underline_runs_is_ignored():
    # 无下划线 run 的裸标签段不能被清空（防误伤普通标签行）
    p = _cover_paragraph([('学院', False)])
    _fill_cover_paragraphs([p], {'college': '测试学院'}, {'学院': 'college'})
    texts = [r.find(qn('w:t')).text for r in p.findall(qn('w:r'))]
    assert texts == ['学院']


def test_strip_foreign_refs_removes_ole_object_keeps_preview_image():
    # OLE 嵌入件（embeddings 部件）不随切片复制 → o:OLEObject 必须剥掉，
    # 否则 Word 报「文件可能已经损坏」；静态预览图 v:imagedata 保留
    from lxml import etree
    from backend.thesis_builder.template_merge import _strip_foreign_refs
    O = 'urn:schemas-microsoft-com:office:office'
    V = 'urn:schemas-microsoft-com:vml'
    R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:o="{O}" xmlns:v="{V}" xmlns:r="{R}">'
        '<w:r><w:object><v:rect><v:imagedata r:id="rId9"/></v:rect>'
        '<o:OLEObject Type="Embed" r:id="rId12"/></w:object></w:r></w:p>')
    el = etree.fromstring(xml)
    _strip_foreign_refs(el)
    assert el.find('.//{%s}OLEObject' % O) is None
    assert el.find('.//{%s}imagedata' % V) is not None
